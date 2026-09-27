"""Checks for the bounded public Texas extraction and candidate distinction."""

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from texas.project import stage
from texas.tpit import extract

TITLE = (
    "JULY BASE FOR TRANSMISSION OWNER SUBMISSION FILE - "
    "TRANSMISSION PROJECT INFORMATION TRACKING (TPIT) FUTURE PROJECTS AS OF 7/13/2026"
)


def _inputs(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PlannedTPIT071326NoCost"
    sheet.append([TITLE] + [None] * 32)
    headings = [None] * 33
    headings[0] = "ERCOT Project Number"
    headings[6] = 'Transmission Status "under construction, planned or conceptual"'
    headings[11] = "Projected In-Service Date (Month/Yr)"
    headings[18] = "County Location for Substation or Starting Point for a Line"
    sheet.append(headings)
    for row_number in range(3, 361):
        row = [None] * 33
        row[0] = row_number
        row[1] = f"Project {row_number}"
        row[6] = "Planned"
        row[8] = "Test owner"
        row[11] = date(2027, 8, 19)
        row[18] = "Williamson"
        if row_number in (67, 285):
            row[5 if row_number == 67 else 4] = "Georgetown East"
            row[19] = "Williamson"
        elif row_number == 105:
            row[4] = "Gabriel Substation"
        elif row_number == 333:
            row[4] = "Georgetown Substation"
        row[32] = "fixture"
        sheet.append(row)
    book_path = tmp_path / "source.xlsx"
    workbook.save(book_path)
    features = []
    for name in ("GEORGETOWN EAST", "GABRIEL", "GEORGETOWN"):
        features.append({"geometry": {"coordinates": [-97.6, 30.6]},
                         "properties": {"Name": name, "GlobalID": name, "OWNER": "LCRA"}})
    gis_path = tmp_path / "gis.json"
    gis_path.write_text(json.dumps({"features": features}))
    return book_path, gis_path


def test_tpit_extract_preserves_month_and_separate_candidate_tier(tmp_path):
    workbook, gis = _inputs(tmp_path)
    rows, summary = extract(workbook, gis)
    assert summary["observations"] == 358
    assert summary["candidate_locations"] == 4
    assert summary["published_projects"] == 0
    assert rows[102]["projected_in_service"] == {
        "raw": "2027-08-19T00:00:00", "value": "2027-08", "precision": "month"
    }
    assert rows[102]["candidate_location"]["independent_review"] is False
    assert rows[64]["candidate_location"]["facility_name"] == "GEORGETOWN EAST"
    assert rows[0]["candidate_location"] is None


def test_tpit_rejects_ambiguous_facility_and_duplicate_id(tmp_path):
    workbook, gis = _inputs(tmp_path)
    data = json.loads(gis.read_text())
    data["features"].append(data["features"][1])
    gis.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="unique named facility"):
        extract(workbook, gis)

    _, gis = _inputs(tmp_path)
    book = load_workbook(workbook)
    book.active["A4"] = book.active["A3"].value
    book.save(workbook)
    with pytest.raises(ValueError, match="ID uniqueness"):
        extract(workbook, gis)


def test_staged_publication_contains_only_four_labeled_candidates():
    folder = Path(__file__).resolve().parents[2] / "data" / "texas"
    rows = json.loads((folder / "planned-observations.json").read_text())
    summary = json.loads((folder / "planned-summary.json").read_text())
    source, projects = stage(rows, summary)
    assert source["project_count"] == 4
    assert {p["native_id"] for p in projects} == {"80546B", "92629", "80546C", "85973"}
    assert all(p["location_review"] == "unreviewed" for p in projects)
    assert all(p["center"]["basis"] == "one" for p in projects)
    assert all(p["in_service"]["precision"] == "month" for p in projects)
    assert len({(p["center"]["lat"], p["center"]["lon"]) for p in projects}) == 3

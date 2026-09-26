from __future__ import annotations

import copy
import io
import json
from pathlib import Path

import openpyxl
import pytest

from verified.build import _activities, _territories, build_snapshot, canonical_json_sha
from verified.fetch import checked_archive
from verified.geography import GeographyIndex
from verified.validate import (
    load_published_snapshot,
    validate_published,
    validate_records_schema,
    validate_schema,
    validate_snapshot,
)

ROOT = Path(__file__).resolve().parents[2]
CACHE = ROOT / "data" / "verified" / "cache" / "f8612024.zip"


def test_published_artifacts_bind_manifest_schemas_and_exact_denominators() -> None:
    validate_published(ROOT)
    manifest = json.loads((ROOT / "data" / "verified" / "manifest.json").read_text(encoding="utf-8"))
    coverage = json.loads((ROOT / "data" / "verified" / "coverage.json").read_text(encoding="utf-8"))
    validate_schema(manifest, ROOT / "data" / "verified" / "schemas" / "manifest.schema.json")
    validate_schema(coverage, ROOT / "data" / "verified" / "schemas" / "coverage.schema.json")
    envelope_schema = ROOT / "data" / "verified" / "schemas" / "record-envelope.schema.json"
    for filename in (
        "sources.json",
        "utilities.json",
        "utility-activities.json",
        "service-territory.json",
        "assertions.json",
        "quarantine.json",
    ):
        validate_schema(json.loads((ROOT / "data" / "verified" / filename).read_text(encoding="utf-8")), envelope_schema)
    assert (ROOT / "data" / "verified" / ".gitattributes").read_text(encoding="utf-8") == "*.json text eol=lf\n"
    assert coverage["counts"] == {
        "assertions": 11_866,
        "comparable_field_conflicts": 0,
        "conflicting_county_rows": 11,
        "independently_corroborated_service_claims": 0,
        "county_identity_quarantine_rows": 48,
        "quarantine": 51,
        "rejected_county_rows": 0,
        "resolved_county_rows": 11_818,
        "service_territory_by_validation_status": {"accepted": 11_818, "conflicting": 11, "unresolved": 37},
        "service_territory_rows": 11_866,
        "sources": 2,
        "unresolved_county_rows": 37,
        "unknown_utility_rows": 3,
        "utilities": 3_413,
        "utilities_by_validation_status": {"accepted": 3_380, "needs_review": 33},
        "utility_activities": 1_706,
    }


@pytest.mark.skipif(not CACHE.exists(), reason="reviewed raw archive intentionally remains in the ignored local cache")
def test_reviewed_archive_replay_is_content_deterministic() -> None:
    snapshot = build_snapshot(ROOT, generated_at="2030-01-01T00:00:00Z")
    published, geography = load_published_snapshot(ROOT)
    assert snapshot["dataset"] == published["dataset"]
    assert snapshot["generated_at"] != published["generated_at"]
    validate_snapshot(snapshot, geography)


def test_cross_record_validator_rejects_false_lineage_count_hash_and_geography() -> None:
    snapshot, geography = load_published_snapshot(ROOT)
    mutations = []
    false_lineage = copy.deepcopy(snapshot)
    false_lineage["assertions"][0]["independently_corroborated"] = True
    mutations.append(false_lineage)
    bad_hash = copy.deepcopy(snapshot)
    bad_hash["utilities"][0]["evidence"]["source_sha256"] = "0" * 64
    mutations.append(bad_hash)
    false_count = copy.deepcopy(snapshot)
    false_count["coverage"]["counts"]["utilities"] += 1
    mutations.append(false_count)
    bad_parent = copy.deepcopy(snapshot)
    accepted = next(record for record in bad_parent["service-territory"] if record["validation_status"] == "accepted")
    accepted["state_fips"] = "99"
    mutations.append(bad_parent)
    for candidate in mutations:
        with pytest.raises(ValueError):
            validate_snapshot(candidate, geography)


def test_unknown_frame_keys_and_ambiguous_counties_stay_quarantined() -> None:
    snapshot, _ = load_published_snapshot(ROOT)
    withheld = [item for item in snapshot["utility-activities"] if item["eia_utility_id"] == "88888"]
    assert len(withheld) == 3
    assert all(item["validation_status"] == "rejected" for item in withheld)
    ambiguous = [item for item in snapshot["service-territory"] if item["validation_status"] == "conflicting"]
    assert {item["county_raw"] for item in ambiguous} >= {"Roanoke", "Fairfax", "Baltimore", "St Louis"}
    quarantined = {item["record_id"] for item in snapshot["quarantine"]}
    assert all(item["id"] in quarantined for item in ambiguous)


def test_county_matcher_preserves_city_ambiguity_and_accepts_only_unique_aliases() -> None:
    geography = {
        "states": [{"state_fips": "51", "usps": "VA"}, {"state_fips": "24", "usps": "MD"}],
        "counties": [
            {"county_geoid": "51161", "state_fips": "51", "state_usps": "VA", "name": "Roanoke", "full_name": "Roanoke County"},
            {"county_geoid": "51770", "state_fips": "51", "state_usps": "VA", "name": "Roanoke", "full_name": "Roanoke city"},
            {
                "county_geoid": "24033",
                "state_fips": "24",
                "state_usps": "MD",
                "name": "Prince George's",
                "full_name": "Prince George's County",
            },
        ],
    }
    index = GeographyIndex(geography)
    assert index.match_county("VA", "Roanoke").status == "conflicting"
    assert index.match_county("VA", "Roanoke City").county_geoid == "51770"
    assert index.match_county("MD", "Prince Georges").county_geoid == "24033"
    assert index.match_county("VA", "Bedford City").status == "unresolved"


def test_cache_is_required_and_never_replaced_by_fixture_data(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        checked_archive(tmp_path / "missing.zip")
    corrupt = tmp_path / "f8612024.zip"
    corrupt.write_bytes(b"not reviewed source bytes")
    with pytest.raises(ValueError, match="hash"):
        checked_archive(corrupt)


def test_normalized_geography_digest_is_line_ending_independent_and_domain_schema_is_strict() -> None:
    value = {"states": [{"state_fips": "13"}], "counties": []}
    lf = json.dumps(value, indent=2) + "\n"
    crlf = lf.replace("\n", "\r\n")
    assert canonical_json_sha(json.loads(lf)) == canonical_json_sha(json.loads(crlf))
    snapshot, _ = load_published_snapshot(ROOT)
    malformed = {**snapshot["utilities"][0], "unexpected": True}
    with pytest.raises(ValueError, match="record"):
        validate_records_schema([malformed], ROOT / "data" / "verified" / "schemas" / "domain-record.schema.json")


def _workbook_bytes(sheets: dict[str, tuple[int, list[object]]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, (header_row, values) in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in range(1, header_row):
            sheet.cell(row, 1, "title")
        sheet.append([f"field-{index}" for index in range(len(values))])
        sheet.append(values)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_same_vintage_name_and_ownership_conflicts_are_rejected_and_quarantined() -> None:
    geography = GeographyIndex(
        {
            "states": [{"state_fips": "13", "usps": "GA"}],
            "counties": [
                {
                    "county_geoid": "13001",
                    "state_fips": "13",
                    "state_usps": "GA",
                    "name": "Appling",
                    "full_name": "Appling County",
                }
            ],
        }
    )
    known = {"1": {"eia_utility_id": "1", "name": "Reference Utility", "ownership": "Investor Owned"}}
    activity = [2024, 1, "Different Utility", "GA", "Municipal", None, *([None] * 26)]
    activity_bytes = _workbook_bytes({"States": (2, activity), "Territories": (2, [])})
    activities, activity_quarantine = _activities(activity_bytes, geography, known)
    assert activities[0]["validation_status"] == "rejected"
    assert activity_quarantine[0]["reason"] == "comparable_field_conflict"
    assert activity_quarantine[0]["conflicting_fields"] == ["name", "ownership"]

    territory = [2024, 1, "Different Utility", None, "GA", "Appling"]
    territory_bytes = _workbook_bytes({"Counties_States": (1, territory), "Counties_Territories": (1, [])})
    territories, assertions, territory_quarantine = _territories(territory_bytes, geography, known)
    assert territories[0]["county_geoid"] == "13001"
    assert territories[0]["validation_status"] == "rejected"
    assert assertions[0]["validation_status"] == "rejected"
    assert territory_quarantine[0]["reason"] == "comparable_field_conflict"

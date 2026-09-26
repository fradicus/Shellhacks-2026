from datetime import datetime

import pytest
from openpyxl import Workbook

import national.iso_ne as iso_ne


def geography():
    return {"states": [{"usps": "MA", "state_fips": "25"}]}


def workbook(path, *, duplicate=False, changed_header=False):
    book = Workbook()
    sheet = book.active
    sheet.title = iso_ne.SHEET
    for column, value in iso_ne.EXPECTED_HEADERS.items():
        sheet.cell(1, column, "Changed" if changed_header and column == 58 else value)
    values = [datetime(2028, 12, 31), 2016, "TBD", datetime(2027, 1, 1)]
    for offset, value in enumerate(values, start=2):
        project_id = 1 if duplicate and offset == 3 else offset - 1
        sheet.cell(offset, 1, "Reliability Upgrade")
        sheet.cell(offset, 2, "1a")
        sheet.cell(offset, 3, project_id)
        sheet.cell(offset, 4, "MA" if offset < 5 else None)
        sheet.cell(offset, 5, "Owner ")
        sheet.cell(offset, 6, "One Company, Another Company" if offset == 2 else None)
        sheet.cell(offset, 8, value)
        if isinstance(value, datetime):
            sheet.cell(offset, 8).number_format = "mm/yyyy"
        sheet.cell(offset, 10, f"Project {project_id}")
        sheet.cell(offset, 58, "Planned")
        sheet.cell(offset, 109, "NR")
    book.save(path)


def test_adapter_preserves_month_year_unknown_and_unsplit_owner(tmp_path, monkeypatch):
    path = tmp_path / "source.xlsx"
    workbook(path)
    monkeypatch.setattr(iso_ne, "EXPECTED_ROWS", 4)
    projects = iso_ne.parse(path, geography(), "a" * 64)
    assert projects[0]["in_service"] == {
        "raw": "2028-12-31T00:00:00",
        "value": "2028-12",
        "precision": "month",
    }
    assert projects[1]["in_service"] == {"raw": "2016", "value": "2016", "precision": "year"}
    assert projects[2]["in_service"] == {"raw": "TBD", "value": None, "precision": "unknown"}
    assert projects[0]["other_owners"] == ["One Company, Another Company"]
    assert projects[3]["states"] == [] and projects[3]["center"] is None


@pytest.mark.parametrize("changed_header,duplicate,match", [(True, False, "column 58"), (False, True, "duplicate")])
def test_adapter_fails_on_header_drift_and_duplicate_ids(tmp_path, monkeypatch, changed_header, duplicate, match):
    path = tmp_path / "source.xlsx"
    workbook(path, duplicate=duplicate, changed_header=changed_header)
    monkeypatch.setattr(iso_ne, "EXPECTED_ROWS", 4)
    with pytest.raises(ValueError, match=match):
        iso_ne.parse(path, geography(), "a" * 64)

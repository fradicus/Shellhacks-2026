"""Rebuild data/fixtures/ from the sponsor workbook. Run: cd pipeline && uv run python ../tests/golden/build_fixtures.py

Deterministic: rerunning produces byte-identical files. The sample is real sponsor data, labeled `historical`.
"""

import hashlib
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "pipeline"))

import openpyxl  # noqa: E402

from common import REPO_ROOT, project_id, validate, write_json  # noqa: E402
from matches.core import center, overlaps, priority_sort  # noqa: E402

WORKBOOK = REPO_ROOT / "docs/Sperry-Tech-Challenge/Projects_Overlaps.xlsx"
FIXTURES = REPO_ROOT / "data/fixtures"
SOURCE_ID = "sperry-sample"
ANALYSIS_DATE = "2026-09-26"  # roadmap analysis_date
UTILITY = {"Dominion Energy South Carolina": "DESC", "Georgia Power": "GPC"}


def parse_date(v) -> str:
    """Workbook dates arrive as 'M/D/YYYY' strings or as datetimes (openpyxl converts serials, e.g. 45809 -> 2025-06-01)."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return datetime.strptime(v.strip(), "%m/%d/%Y").date().isoformat()


def raw_date(v) -> str:
    return v.strip() if isinstance(v, str) else f"{v.month}/{v.day}/{v.year}"


def read_workbook():
    wb = openpyxl.load_workbook(WORKBOOK)  # formulas stay unevaluated; centers are recomputed by core.center
    rows = list(wb["projects"].iter_rows(values_only=True))
    head = rows[0]
    projects = []
    for n, r in enumerate(rows[1:], start=2):
        if not r[0]:
            continue
        d = dict(zip(head, r, strict=True))
        d["_row"] = n
        projects.append(d)
    ov_rows = list(wb["overlaps"].iter_rows(values_only=True))
    overlaps_sheet = [dict(zip(ov_rows[0], r, strict=True)) for r in ov_rows[1:] if r[0]]
    return projects, overlaps_sheet


def golden_projects(rows):
    out = []
    for d in rows:
        out.append({
            "project_id": d["project_id"],
            "utility": UTILITY[d["utility"]],
            "state": d["state"],
            "project_name": d["project_name"],
            "endpoints": [
                {"name": d["name_a"], "lat": d["lat_a"], "lon": d["lon_a"]},
                {"name": d["name_b"], "lat": d["lat_b"], "lon": d["lon_b"]},
            ],
            "in_service_raw": raw_date(d["in_service_date"]),
            "in_service_date": parse_date(d["in_service_date"]),
            "row": d["_row"],
        })
    return out


def golden_overlaps(sheet):
    return [{
        "overlap_id": r["overlap_id"],
        "distance_mi": r["distance_mi"],
        "time_gap_days": r["time_gap (day)"],
        "project_id_a": r["project_id_a"],
        "project_id_b": r["project_id_b"],
    } for r in sheet]


def ui_fixtures(golden):
    sha = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    sources = [{
        "_id": SOURCE_ID,
        "publisher": "Sperry Tech (ShellHacks 2026 Gridlock challenge)",
        "title": "Projects_Overlaps.xlsx: sponsor sample of 10 projects and 6 overlaps",
        "sha256": sha,
        "pages": None,
        "public_status": "public",
        "local_path": "docs/Sperry-Tech-Challenge/Projects_Overlaps.xlsx",
    }]
    projects, locations = [], []
    for g in golden:
        key = f"{g['utility']}:{g['project_id']}"
        eps = []
        for i, e in enumerate(g["endpoints"]):
            cell = "F:G" if i == 0 else "I:J"
            loc = {"_id": f"{key}#{i}", "project_key": key, "endpoint_index": i, "name": e["name"]}
            if e["lat"] is None or e["lon"] is None:
                loc |= {"confidence": "rejected", "evidence": f"No coordinates in the sponsor sample (projects row {g['row']})"}
            else:
                loc |= {"confidence": "high", "lat": e["lat"], "lon": e["lon"],
                        "evidence": f"Coordinates from the sponsor sample workbook, projects row {g['row']} cols {cell}"}
            eps.append(loc)
        locations += eps
        c = center(eps)
        projects.append({
            "_id": project_id(key, SOURCE_ID),
            "project_key": key,
            "utility": g["utility"],
            "owner_code": None,
            "owner_basis": "utility column of the sponsor sample",
            "native_id": g["project_id"],
            "name": g["project_name"],
            "state": g["state"],
            "source": {"source_id": SOURCE_ID, "page": None},
            "in_service": {"raw": g["in_service_raw"], "date": g["in_service_date"], "precision": "day"},
            "active": True,
            "center": c,
            "location_confidence": "high" if c else None,
            "geo": {"type": "Point", "coordinates": [c["lon"], c["lat"]]} if c else None,
        })
    matches = [m | {"review_state": "needs_review"} for m in priority_sort(overlaps(projects, ANALYSIS_DATE))]
    version_changes = [{
        "_id": "DESC:0139 M,N|in_service.date|desc-2024-2028>desc-2025-2029",
        "project_key": "DESC:0139 M,N",
        "from_source": "desc-2024-2028",
        "to_source": "desc-2025-2029",
        "field": "in_service.date",
        "old": "2024-12-31",
        "new": "2026-05-31",
        "from_page": 3,
        "to_page": 2,
    }]
    return {
        "sources": sources, "projects": projects, "locations": locations, "matches": matches,
        "version_changes": version_changes, "briefs": [], "extractions": [], "reviews": [], "runs": [], "coverage": [],
    }


SCHEMA_OF = {
    "sources": "source", "projects": "project", "locations": "location", "matches": "match", "version_changes": "version_change",
    "briefs": "brief", "extractions": "extraction", "reviews": "review", "runs": "run", "coverage": "coverage",
}


def main():
    rows, sheet = read_workbook()
    golden = golden_projects(rows)
    write_json(FIXTURES / "golden/projects.json", golden)
    write_json(FIXTURES / "golden/overlaps.json", golden_overlaps(sheet))
    for name, records in ui_fixtures(golden).items():
        for rec in records:
            validate(rec, SCHEMA_OF[name])
        write_json(FIXTURES / f"{name}.json", records)
    print(f"wrote {FIXTURES}")


if __name__ == "__main__":
    main()

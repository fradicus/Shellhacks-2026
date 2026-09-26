"""Audit the sponsor XLSX against independent math, committed fixtures and live core.

Run from pipeline: uv run python ../tests/e2e/golden_independent.py
Add --report ../reports/qa/golden.md to save the same Markdown printed to stdout.
This reference imports no production matcher code. core_snapshot.py is the isolated
system-under-test adapter, invoked only after the reference has been calculated.
"""

import argparse
import hashlib
import json
import math
import subprocess
import sys
from datetime import date, datetime
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "docs/Sperry-Tech-Challenge/Projects_Overlaps.xlsx"
UTILITIES = {"Dominion Energy South Carolina": "DESC", "Georgia Power": "GPC"}


def day(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return from_excel(value, epoch).date()
    return datetime.strptime(value, "%m/%d/%Y").date()


def rows(sheet):
    values = sheet.iter_rows(values_only=True)
    headers = next(values)
    return [dict(zip(headers, row, strict=True)) for row in values if any(v is not None for v in row)]


def distance(a, b):
    p, q = math.radians(a["lat"]), math.radians(b["lat"])
    h = math.sin((q - p) / 2) ** 2 + math.cos(p) * math.cos(q) * math.sin(math.radians(b["lon"] - a["lon"]) / 2) ** 2
    h = min(1, max(0, h))
    return 2 * 3958.8 * math.atan2(math.sqrt(h), math.sqrt(1 - h))


def reference(path=WORKBOOK):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        source, sheet = rows(book["projects"]), rows(book["overlaps"])
        projects = []
        centers = {}
        for row in source:
            utility = UTILITIES[row["utility"]]
            key = f"{utility}:{row['project_id']}"
            endpoints = [{"name": row[f"name_{e}"], "lat": row[f"lat_{e}"], "lon": row[f"lon_{e}"]} for e in "ab"]
            located = [e for e in endpoints if e["lat"] is not None and e["lon"] is not None]
            centers[key] = (
                {
                    "lat": sum(e["lat"] for e in located) / len(located),
                    "lon": sum(e["lon"] for e in located) / len(located),
                    "basis": "one" if len(located) == 1 else "two",
                }
                if located
                else None
            )
            projects.append(
                {
                    "project_key": key,
                    "utility": utility,
                    "endpoints": endpoints,
                    "in_service": {"date": day(row["in_service_date"], book.epoch).isoformat(), "precision": "day"},
                }
            )
    finally:
        book.close()
    pairs = {}
    for a, b in combinations(projects, 2):
        if a["utility"] == b["utility"]:
            continue
        ka, kb = sorted((a["project_key"], b["project_key"]))
        ca, cb = centers[ka], centers[kb]
        pairs[f"{ka}__{kb}"] = {
            "distance_mi": distance(ca, cb) if ca and cb else None,
            "time_gap_days": abs(
                (date.fromisoformat(a["in_service"]["date"]) - date.fromisoformat(b["in_service"]["date"])).days
            ),
        }
    return source, sheet, projects, centers, pairs


def audit(source, sheet, projects, centers, pairs, core, fixtures):
    errors = []
    notes = []

    def check(ok, message):
        if not ok:
            errors.append(message)

    def compare_float(label, actual, expected):
        if actual is None or expected is None:
            check(actual == expected, f"{label}: {actual!r} != {expected!r}")
            return
        delta = actual - expected
        if delta != 0:
            notes.append(f"{label}: actual={actual!r}, reference={expected!r}, delta={delta!r}")
        check(math.isclose(actual, expected, rel_tol=0, abs_tol=1e-10), f"{label}: delta={delta!r}")

    check(len(projects) == 10 and len(centers) == 10, "Expected ten unique sponsor projects")
    check(len(pairs) == 25, "Expected 25 cross-utility pairs")
    overlaps = {key: p for key, p in pairs.items() if p["distance_mi"] is not None and p["distance_mi"] < 25}
    expected = {}
    for row in sheet:
        keys = [f"{UTILITIES[row[f'utility_{side}']]}:{row[f'project_id_{side}']}" for side in "ab"]
        expected["__".join(sorted(keys))] = row
    check(len(sheet) == len(expected) == 6, "Expected six unique workbook overlaps")
    check({r["overlap_id"] for r in sheet} == {f"OVL_{i}" for i in range(1, 7)}, "Workbook overlap IDs differ")
    check(set(overlaps) == set(expected), "Independent overlap set differs from workbook")
    check(len(pairs) - len(overlaps) == 19, "Expected nineteen excluded cross-utility pairs")
    actual = {m["_id"]: m for m in core["matches"]}
    check(len(actual) == len(core["matches"]), "Duplicate core match IDs")
    check(set(actual) == set(overlaps), "Core overlap set differs from independent result")
    check(set(core["centers"]) == set(centers), "Core project keys differ")
    fixture_by_key = {f"{p['utility']}:{p['project_id']}": p for p in fixtures}
    check(len(fixtures) == len(fixture_by_key) == 10 and set(fixture_by_key) == set(centers), "Fixture project keys differ")

    for row, project in zip(source, projects, strict=True):
        key = project["project_key"]
        center = centers[key]
        for axis in ("lat", "lon"):
            compare_float(f"{key} workbook center {axis}", row[f"{axis}_center"], center[axis] if center else None)
            compare_float(
                f"{key} core center {axis}", (core["centers"].get(key) or {}).get(axis), center[axis] if center else None
            )
        check((core["centers"].get(key) or {}).get("basis") == (center or {}).get("basis"), f"{key} center basis differs")
        fixture = fixture_by_key.get(key, {})
        check(fixture.get("endpoints") == project["endpoints"], f"{key} fixture endpoints differ")
        check(fixture.get("in_service_date") == project["in_service"]["date"], f"{key} fixture date differs")
        neighbors = {k.split("__")[1] if k.split("__")[0] == key else k.split("__")[0] for k in overlaps if key in k.split("__")}
        listed = {str(row[f"overlap_{i}"]) for i in range(1, 4) if row[f"overlap_{i}"]}
        check(
            {n.split(":", 1)[1] for n in neighbors} == listed and len(neighbors) == row["overlap_count"],
            f"{key} project-sheet overlap annotations differ",
        )

    table = [
        "| Pair | Projects | Reference mi | Workbook mi | Raw minus workbook | Core minus reference | Gap days |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    for key, row in sorted(expected.items(), key=lambda item: item[1]["overlap_id"]):
        ref = overlaps.get(key)
        if ref is None:
            continue
        dist, gap = ref["distance_mi"], ref["time_gap_days"]
        check(round(dist, 2) == row["distance_mi"], f"{key} workbook distance differs at two decimals")
        check(gap == row["time_gap (day)"], f"{key} workbook gap differs")
        result = actual.get(key, {})
        compare_float(f"{key} core distance", result.get("distance_mi"), dist)
        check(result.get("time_gap_days") == gap, f"{key} core gap differs")
        check(result.get("band") == (0 if dist < 10 else 1), f"{key} core band differs")
        check([result.get("a"), result.get("b")] == key.split("__"), f"{key} core pair members differ")
        check(result.get("view") == "historical", f"{key} core view differs")
        delta = result["distance_mi"] - dist if result.get("distance_mi") is not None else None
        table.append(
            f"| {row['overlap_id']} | {key} | {dist!r} | {row['distance_mi']:.2f} | "
            f"{dist - row['distance_mi']!r} | {delta!r} | {gap} |"
        )
    order = sorted(
        overlaps,
        key=lambda k: (0 if overlaps[k]["distance_mi"] < 10 else 1, overlaps[k]["time_gap_days"], overlaps[k]["distance_mi"], k),
    )
    check([m["_id"] for m in core["matches"]] == order, "Core priority order differs")
    check([m.get("rank") for m in core["matches"]] == list(range(1, len(order) + 1)), "Core ranks differ")
    ranked_ids = [expected[k]["overlap_id"] for k in order if k in expected]
    check(ranked_ids == ["OVL_2", "OVL_3", "OVL_1", "OVL_4", "OVL_5", "OVL_6"], "Sponsor priority order differs")
    return errors, notes, table, ranked_ids


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    source, sheet, projects, centers, pairs = reference()
    snapshot = subprocess.run(
        [sys.executable, str(Path(__file__).with_name("core_snapshot.py"))],
        input=json.dumps(projects),
        text=True,
        capture_output=True,
        check=True,
    )
    core = json.loads(snapshot.stdout)
    fixtures = json.loads((ROOT / "data/fixtures/golden/projects.json").read_text())
    errors, notes, table, order = audit(source, sheet, projects, centers, pairs, core, fixtures)
    report = "\n".join(
        [
            "# Independent golden verification",
            "",
            f"Source: `{WORKBOOK.relative_to(ROOT)}` (`projects!A1:Q11`, `overlaps!A1:I7`).",
            f"SHA-256: `{hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()}`.",
            "",
            "Reference reads original endpoints and mixed Excel/text dates with openpyxl, then computes arithmetic centers,",
            "haversine with R=3958.8 mi (atan2 form), strict distance <25 mi, exact date gaps and priority independently.",
            "Live production core is called in a separate process using the same workbook inputs. Committed golden",
            "fixture endpoints and dates are checked against those inputs. No production math is used by the reference.",
            "",
            "Workbook distances are published to two decimals; raw rounding deltas below are expected, not hidden.",
            "Every nonzero core/center delta is listed. Absolute tolerance for floating-point comparison is 1e-10;",
            "pair membership, gaps, bands, ranks and workbook two-decimal distances must agree exactly.",
            "",
            *table,
            "",
            f"Priority: {', '.join(order)}.",
            f"Checked {len(projects)} projects and {len(pairs)} cross-utility pairs; "
            "expected six overlaps and nineteen exclusions.",
            "",
            "## Nonzero center/core differences",
            "",
            *(notes or ["None."]),
            "",
            "## Result",
            "",
            *(errors or ["PASS: workbook, independent calculation, fixtures and live core agree within the stated precision."]),
            "",
        ]
    )
    print(report)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(report)
    return int(bool(errors))


if __name__ == "__main__":
    sys.exit(main())

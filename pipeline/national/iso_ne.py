"""Deterministic ISO-NE RSP workbook adapter for its reviewed sortable sheet."""

from __future__ import annotations

import re
import warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SOURCE_ID = "iso-ne-rsp-2026-06"
SHEET = "RSP_sortable"
EXPECTED_ROWS = 1_024
EXPECTED_HEADERS = {
    1: "Primary Driver",
    2: "Part#",
    3: "Project ID",
    4: "State",
    5: "Primary Equipment Owner",
    6: "Other Equipment Owner(s)",
    7: "Footnote Number",
    8: "Projected In-Service Month/Year",
    9: "Major Project",
    10: "Project",
    58: "Jun-26 Status",
    109: "Jun-26 Estimated PTF Costs",
}
STATUS_GROUPS = {
    "planned": "planned",
    "under construction": "under_construction",
    "proposed": "proposed",
    "in-service": "in_service",
    "cancelled": "cancelled",
}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _raw(value: Any) -> str | int | float | bool | None:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _milestone(value: Any) -> dict[str, str | None]:
    if isinstance(value, datetime | date):
        return {"raw": value.isoformat(), "value": f"{value.year:04d}-{value.month:02d}", "precision": "month"}
    if type(value) is int and 1900 <= value <= 2200:
        return {"raw": str(value), "value": str(value), "precision": "year"}
    if isinstance(value, float) and value.is_integer() and 1900 <= value <= 2200:
        year = int(value)
        return {"raw": str(year), "value": str(year), "precision": "year"}
    if _text(value) in (None, "TBD"):
        return {"raw": _text(value), "value": None, "precision": "unknown"}
    raise ValueError(f"unsupported ISO-NE milestone {value!r}")


def _project_id(value: Any, row: int) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int | float) or int(value) != value:
        raise ValueError(f"{SHEET} row {row}: non-integral Project ID {value!r}")
    return str(int(value))


def _workbook_rows(path: Path) -> tuple[list[Any], list[tuple[Any, ...]]]:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Cannot parse header or footer.*", category=UserWarning)
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if SHEET not in workbook.sheetnames:
            raise ValueError(f"ISO-NE workbook is missing {SHEET!r}")
        sheet = workbook[SHEET]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        workbook.close()
    return headers, rows


def parse(path: Path, geography: dict[str, Any], source_sha256: str) -> list[dict[str, Any]]:
    headers, rows = _workbook_rows(path)
    for column, expected in EXPECTED_HEADERS.items():
        actual = _header(headers[column - 1] if column <= len(headers) else None)
        if actual != expected:
            raise ValueError(f"{SHEET} column {column}: expected {expected!r}, found {actual!r}")
    state_fips = {state["usps"]: state["state_fips"] for state in geography["states"]}
    projects = []
    seen: set[str] = set()
    for row_number, cells in enumerate(rows, start=2):
        native_id = _project_id(cells[2] if len(cells) > 2 else None, row_number)
        if native_id is None:
            continue
        if native_id in seen:
            raise ValueError(f"{SHEET} row {row_number}: duplicate Project ID {native_id}")
        seen.add(native_id)
        state_raw = _text(cells[3])
        if state_raw is not None and state_raw not in state_fips:
            raise ValueError(f"{SHEET} row {row_number}: unknown State {state_raw!r}")
        status = _text(cells[57])
        status_key = status.lower() if status else ""
        if status_key not in STATUS_GROUPS:
            raise ValueError(f"{SHEET} row {row_number}: unsupported current status {status!r}")
        name = _text(cells[9])
        owner = _text(cells[4])
        if not name or not owner:
            raise ValueError(f"{SHEET} row {row_number}: project name and primary owner are required")
        other_owner = _text(cells[5])
        cost_raw = cells[108]
        cost_usd = cost_raw if type(cost_raw) in (int, float) else None
        project = {
            "_id": f"iso-ne:{native_id}",
            "source_id": SOURCE_ID,
            "native_id": native_id,
            "name": name,
            "description": None,
            "owner": owner,
            "other_owners": [other_owner] if other_owner else [],
            "planning_region": "iso-ne",
            "states": [state_fips[state_raw]] if state_raw else [],
            "counties": [],
            "geography_basis": "source_state" if state_raw else None,
            "status": status,
            "status_group": STATUS_GROUPS[status_key],
            "in_service": _milestone(cells[7]),
            "center": None,
            "location_review": "unlocated",
            "primary_driver": _text(cells[0]),
            "part": _text(cells[1]),
            "major_project": _text(cells[8]),
            "footnote_number": _raw(cells[6]),
            "estimated_ptf_cost": {"raw": _raw(cost_raw), "usd": cost_usd},
            "evidence": {
                "page": None,
                "sheet": SHEET,
                "row": row_number,
                "source_sha256": source_sha256,
                "raw": {
                    "Project ID": _raw(cells[2]),
                    "State": _raw(cells[3]),
                    "Primary Equipment Owner": _raw(cells[4]),
                    "Other Equipment Owner(s)": _raw(cells[5]),
                    "Projected In-Service Month/Year": _raw(cells[7]),
                    "Major Project": _raw(cells[8]),
                    "Project": _raw(cells[9]),
                    "Jun-26 Status": _raw(cells[57]),
                    "Jun-26 Estimated PTF Costs": _raw(cells[108]),
                },
            },
        }
        projects.append(project)
    if len(projects) != EXPECTED_ROWS:
        raise ValueError(f"{SHEET}: expected {EXPECTED_ROWS} projects, found {len(projects)}")
    return sorted(projects, key=lambda project: int(project["native_id"]))

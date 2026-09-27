"""MISO South and SPP planning workbooks for LA, AR, MS (and MISO's KY rows): C40 OSM name candidates.

From pipeline/:
  uv run python -m southeast.misospp fetch --cache <dir>   # network: 2 MISO + 10 SPP workbooks, OSM for 4 states
  uv run python -m southeast.misospp build --cache <dir> [--check]
MISO: MTEP25 Appendix A (project and facility sheets) and the MTEP26 projects-under-evaluation list; one MTEP
project ID is one project. SPP: the newest Quarterly Project Tracking Appendix 1 (3Q 2026) plus each year's fourth
quarter back to 2017 (older years are .xls, which the installed stack does not read); older editions only add
upgrades the newer ones dropped, and keep every edition's changed in-service date as history.
OSM data (c) OpenStreetMap contributors, ODbL 1.0.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import warnings
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import openpyxl

from california.caiso import match
from common import REPO_ROOT, load_json, write_json
from common.names import norm_name
from greatlakes.match import candidate_center, facility_key, voltages_kv
from greatlakes.miso import STATUS
from greatlakes.shared import fetch_into, fetch_osm, osm_extract, verify_cache

from .dense import OSM_DATASET, SE_STATES, locate, slug, write_batch

BATCH = "misospp"
STATES = ("LA", "AR", "MS", "KY")
MISO_A = "miso-mtep25-appendix-a"
MISO_EVAL = "miso-mtep26-eval"
MISO_FILES = {
    MISO_A: ("https://cdn.misoenergy.org/MTEP25%20Appendix%20A%20-%20New%20Local%20Reliability%20Projects720399.xlsx",
             "MTEP25 Appendix A, recommended new projects and facility details (as of 8/28/2025)"),
    MISO_EVAL: ("https://cdn.misoenergy.org/MTEP%20Projects%20Under%20Evaluation368757.xlsx",
                "MTEP Projects Under Evaluation (MTEP26 cycle)"),
}
SPP_INDEX = "https://www.spp.org/spp-documents-filings/?id=18641"
# Newest first. Paths exactly as SPP's document index links them.
SPP_EDITIONS = {
    "2026q3": "/Documents/77416/3Q 2026 Quarterly Project Tracking Report Appendix 1&2.zip",
    "2025q4": "/Documents/75201/4Q 2025 Quarterly Project Tracking Report Appendix 1&2.zip",
    "2024q4": "/Documents/72785/4Q 2024 Quarterly Project Tracking Report Appendix 1& 2.zip",
    "2023q4": "/Documents/70609/4Q 2023 Quarterly Project Tracking Report Appendix 1 & 2.zip",
    "2022q4": "/Documents/68142/4Q 2022 Report Quarterly Project Tracking Report Appendix 1.xlsm",
    "2021q4": "/Documents/65840/Q4 2021 Quarterly Project Tracking Appendix 1.xlsx",
    "2020q4": "/Documents/63259/Q4 2020 Quarterly Project Tracking Appendix 1.xlsx",
    "2019q4": "/Documents/60925/Q4 2019 Appendix 1.xlsx",
    "2018q4": "/Documents/58950/Q4 2018 Appendix 1.xlsx",
    "2017q4": "/Documents/55027/Q4 2017 Quarterly Project Tracking - Appendix 1.xlsx",
}


def spp_file(edition: str) -> str:
    return f"spp-qpt-{edition}{Path(SPP_EDITIONS[edition]).suffix}"


def fetch(cache: Path) -> None:
    cache.mkdir(parents=True, exist_ok=True)
    manifest = load_json(cache / "manifest.json") if (cache / "manifest.json").exists() else {}
    wanted = {f"{s}.xlsx": url for s, (url, _) in MISO_FILES.items()}
    wanted |= {spp_file(e): "https://www.spp.org" + quote(p) for e, p in SPP_EDITIONS.items()}
    for name, url in wanted.items():
        if name not in manifest:
            fetch_into(cache, name, url, manifest)
            write_json(cache / "manifest.json", manifest)
    for state in STATES:
        if f"osm-{state.lower()}.json" not in manifest:
            fetch_osm(cache, state, manifest)
            write_json(cache / "manifest.json", manifest)


PLANNING = {"miso": "Midcontinent Independent System Operator (MISO)", "spp": "Southwest Power Pool (SPP)"}
ACCESS = {"miso": "Public workbook on cdn.misoenergy.org; no login; no CEII marking in the workbook.",
          "spp": "Public workbook linked from spp.org's Quarterly Project Tracking index; no login; no CEII marking."}
FIPS = SE_STATES | {"TX": "48", "OK": "40", "KS": "20", "MO": "29", "NM": "35", "NE": "31", "ND": "38", "SD": "46",
                    "IA": "19", "MT": "30", "WY": "56", "IL": "17", "MN": "27", "IN": "18", "MI": "26", "WI": "55"}
COOP_ENERGY = ["COOPERATIVE ENERGY", "SOUTH MISSISSIPPI", "SMEPA"]
# OSM operator-tag fragments per reported owner: SPP owner codes exactly, MISO submitting-TO names by fragment.
SPP_KEYS = {"AEP": ["SOUTHWESTERN ELECTRIC", "SWEPCO", "AMERICAN ELECTRIC POWER"], "OGE": ["OKLAHOMA GAS", "OG&E"],
            "AECC": ["ARKANSAS ELECTRIC COOP"], "SWPA": ["SOUTHWESTERN POWER ADMIN"], "GRDA": ["GRAND RIVER DAM"],
            "EDE": ["EMPIRE DISTRICT", "LIBERTY"], "WFEC": ["WESTERN FARMERS"], "ETEC": ["EAST TEXAS ELECTRIC"]}
MISO_KEYS = [("ENTERGY", ["ENTERGY"]), ("CLECO", ["CLECO"]), ("ARKANSAS ELECTRIC COOP", ["ARKANSAS ELECTRIC COOP"]),
             ("SOUTH MISSISSIPPI", COOP_ENERGY), ("COOPERATIVE ENERGY", COOP_ENERGY), ("BIG RIVERS", ["BIG RIVERS"]),
             ("1803", ["1803"]), ("GRIDLIANCE", ["GRIDLIANCE"]), ("MISSISSIPPI POWER", ["MISSISSIPPI POWER"])]
# SPP tags some upgrades with a state the owner does not serve (NWE "Groton to Aberdeen" and ITCGP "Thistle" as AR).
# Owners whose transmission system covers the tagged F39 state are accepted as tagged; owners whose system lies
# outside it are excluded; every other owner (a border utility such as OG&E around Fort Smith, or TBD) is accepted
# only when a facility the row names matches an OSM substation in the tagged state.
SPP_CORE = {"AEP": {"LA", "AR"}, "AECC": {"AR"}}  # SWEPCO serves NW Louisiana and SW Arkansas; AECC is Arkansas's G&T
SPP_OUTSIDE = {"SPS": "Texas/New Mexico", "NWE": "South Dakota", "ITCGP": "Kansas", "EKC": "Kansas"}
SPP_GROUP = [("COMPLETE", "in_service"), ("CLOSED OUT", "in_service"), ("IN SERVICE", "in_service"),
             ("WITHDRAWN", "cancelled"), ("IDENTIFIED", "proposed"), ("PENDING", "proposed"), ("ON SCHEDULE", "planned"),
             ("DELAY", "planned"), ("NTC", "planned")]
SPP_REQUIRED = {"uid", "projectowner", "state(s)", "projectname", "upgradename", "projectstatus", "frombusname",
                "tobusname", "projectownerindicatedin-servicedate"}
JUNK_SUBS = {"MULTIPLE", "TBD", "NEW", "N/A", "NA", "VARIOUS", "NONE"}
BUS_KV = re.compile(r"\s+\d+(?:\.\d+)?\s*(?:KV)?$", re.I)
# Name forms the shared parser does not read, rewritten before parsing (the OSM match itself stays exact):
# leading queue/customer tags ("J2143 ", "DEMCO "), "X 161kV: Build ...", "500-230 kV", "Ft", "- Phase 2".
LEAD_TAG = re.compile(r"^(?:(?:[JE]\d{4}|MPFCA|SLEMCO|DEMCO|1803/LAGT(?:/NELPCO)?)\b[\s&/,]*)+")
AT_SITE = re.compile(r"\bat ([A-Z][\w.'’]*(?: [A-Z][\w.'’]*){0,3})(?: \d+(?:/\d+)? ?kV)?$")
PROGRAM = re.compile(r"\bprogram\b|\basset renewal\b", re.I)


def clean(value) -> object:
    return value.isoformat() if isinstance(value, datetime) else value


def day(value) -> str | None:
    """A workbook date cell (datetime or m/d/yyyy text) as YYYY-MM-DD; anything else is unknown."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.year >= 2000 else None
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%m/%d/%Y").date().isoformat()
        except ValueError:
            return None
    return None


def spp_group(status: str) -> str:
    text = " ".join(status.split()).upper()
    return next((group for prefix, group in SPP_GROUP if text.startswith(prefix)), "unknown")


def spp_states(tag) -> list[str]:
    return list(dict.fromkeys(t.strip() for t in str(tag or "").split("/") if t.strip() in FIPS))


def miso_keys(owner: str) -> list[str]:
    up = owner.upper()
    return next((keys for fragment, keys in MISO_KEYS if fragment in up), [])


def spp_rule(owner: str, states: list[str]) -> tuple[str, str | None]:
    """('core'|'border'|'excluded', reason) for an SPP owner and its tagged F39 states."""
    f39 = [s for s in states if s in STATES]
    if SPP_CORE.get(owner, set()) & set(f39):
        return "core", None
    if owner in SPP_OUTSIDE:
        return "excluded", f"owner {owner}'s system is in {SPP_OUTSIDE[owner]}, not the tagged {'/'.join(f39)}"
    return "border", None


def parse_text(name: str) -> str:
    text = LEAD_TAG.sub("", " ".join(name.split()))
    text = re.sub(r"(\d)-(\d)", r"\1/\2", text.replace(":", " "))
    text = re.sub(r"\bFt\.?(?= )", "Fort", text, flags=re.I)
    return re.sub(r"\s*[–-]\s*Phase \d+$", "", " ".join(text.split()))


def locate_text(name: str, description: str | None, facilities: list[dict], keys: list[str]
                ) -> tuple[dict | None, dict]:
    """dense.locate on the rewritten name; "... at <Name> 230 kV" names one site."""
    text = parse_text(name)
    if m := AT_SITE.search(text):
        return locate_names("site", [m[1]], voltages_kv(text), facilities, keys, "name_at_site")
    return locate(text, description, facilities, keys)


def locate_names(kind: str, names: list[str], kv: set[int], facilities: list[dict], keys: list[str],
                 names_from: str) -> tuple[dict | None, dict]:
    """dense.locate for facility names the source states in separate columns (same C40 tiers and C38 guard)."""
    matches = [match(n, facilities, keys, kv) for n in names]
    center = candidate_center(kind, matches)
    found = [m for m in matches if m["status"] == "matched"]
    tier = None
    if center:
        tier = "candidate" if all(m["corroboration"] != ["unique_in_state"] for m in found) else "candidate_unique_name"
    fields = ("id", "name", "operator", "voltage", "lat", "lon")
    endpoints = [{k: v for k, v in m.items() if k not in ("facility", "facility_ids")}
                 | ({"facility": {f: m["facility"].get(f) for f in fields}} if m["status"] == "matched" else {})
                 for m in matches]
    return center, {"rule": "C40", "tier": tier, "independent_review": False, "kind": kind, "reason": None,
                    "names_from": names_from, "voltages_kv": sorted(kv), "operator_keys": keys,
                    "endpoints": endpoints, "dataset": OSM_DATASET}


def facility_names(facilities: list[dict]) -> tuple[str, list[str]] | None:
    """MISO facility rows -> one site, or a line whose From/To both name the only two substations; else None."""
    subs: dict[str, str] = {}
    for f in facilities:
        for name in (f["From Sub"], f["To Sub"]):
            if name and str(name).strip() and norm_name(str(name)) not in JUNK_SUBS:
                subs.setdefault(facility_key(str(name)), " ".join(str(name).split()))
    lines = [f for f in facilities if str(f["Facility Type"] or "").startswith("Line")]
    if len(subs) == 1:
        return ("line" if lines else "site"), list(subs.values())
    if len(subs) == 2 and any({facility_key(str(f["From Sub"] or "")), facility_key(str(f["To Sub"] or ""))}
                              == set(subs) for f in lines):
        return "line", list(subs.values())
    return None


def bus_names(upgrade: str, buses: list) -> list[str]:
    """SPP bus names the upgrade name itself repeats; model bus names that disagree with it are not evidence."""
    text, out = f" {facility_key(upgrade)} ", {}
    for bus in buses:
        name = BUS_KV.sub("", " ".join(str(bus or "").split()))
        key = facility_key(name) if name else ""
        if key and f" {key} " in text:
            out.setdefault(key, parse_text(name))
    return list(out.values())


def event(pid: str, native: str, kind: str, date: str | None, text: str, evidence: dict, facts: str,
          suffix: str) -> dict:
    return {"id": f"{pid}:{suffix}", "type": kind, "date": date, "precision": "day" if date else "unknown",
            "native_project_link": native, "description": text, "evidence": [evidence | {"facts": facts}]}


def history(pid: str, native: str, observations: list[dict]) -> list[dict]:
    """Oldest-first observations -> one planned_milestone per changed expected date, then in_service if reported."""
    events, last = [], None
    for obs in observations[:-1] if observations[-1]["group"] == "in_service" else observations:
        if obs["group"] == "in_service" or not obs["date"] or obs["date"] == last:
            continue
        last = obs["date"]
        events.append(event(pid, native, "planned_milestone", obs["date"],
                            f"Expected in-service {obs['date']}, reported in {obs['title']}.", obs["evidence"],
                            f"{obs['date_field']} = {obs['date']}", f"{obs['edition']}-row-{obs['row']}-expected"))
    newest = observations[-1]
    if newest["group"] == "in_service":
        reported = newest["date"] if newest["date"] and newest["date"] <= newest["evidence"]["retrieved_at"][:10] \
            else None
        events.append(event(pid, native, "in_service", reported,
                            f"Status “{newest['status']}” in {newest['title']}"
                            + (f"; {newest['date_field']} lists {reported}." if reported else
                               "; the source gives no past in-service date for it."),
                            newest["evidence"], f"status = {newest['status']}; {newest['date_field']} = "
                                                f"{newest['date'] or 'blank'}", "in-service"))
    return events


def workbook(data) -> openpyxl.Workbook:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # openpyxl: unsupported data-validation extension
        return openpyxl.load_workbook(data, read_only=True, data_only=True)


def sheet_rows(ws, header_at: int) -> list[tuple[int, dict]]:
    table = list(ws.iter_rows(values_only=True))
    header = [" ".join(str(h).split()) if h is not None else "" for h in table[header_at]]
    return [(n, dict(zip(header, r, strict=False))) for n, r in enumerate(table[header_at + 1:], start=header_at + 2)
            if any(v is not None for v in r)]


def mtep_id(value) -> str:
    # One Appendix A ID cell is stored as a date-formatted number; its serial is the ID.
    return str((value - datetime(1899, 12, 30)).days) if isinstance(value, datetime) else str(value).strip()


def artifact_of(manifest: dict, name: str) -> dict:
    return manifest[name] | {"file": name}


def evidence_of(kind: str, artifact: dict, locator: str, source_date: str | None) -> dict:
    return {"publisher": PLANNING[kind], "url": artifact["url"], "artifact_sha256": artifact["sha256"],
            "locator": locator, "source_date": source_date, "retrieved_at": artifact["retrieved_at"],
            "access_review": ACCESS[kind], "facts": "-"}


def load_osm(cache: Path) -> dict[str, list[dict]]:
    return {s: osm_extract(json.loads((cache / f"osm-{s.lower()}.json").read_bytes()), s) for s in STATES}


def miso_observations(cache: Path, manifest: dict) -> list[dict]:
    """Every MISO project row (MTEP25 Appendix A/B sheets, then the MTEP26 list), oldest file first."""
    out = []
    art = artifact_of(manifest, f"{MISO_A}.xlsx")
    wb = workbook(cache / art["file"])
    sheets = {ws.title: ws for ws in wb.worksheets}
    title = next(sheets["MTEP25 Data Pull - App. A"].iter_rows(values_only=True, max_row=1))[0]
    as_of = datetime.strptime(re.search(r"as of (\d+/\d+/\d{4})", title)[1], "%m/%d/%Y").date().isoformat()
    facilities: dict[str, list[dict]] = {}
    for n, f in sheet_rows(sheets["MTEP25 Data Pull - Facility"], 1):
        f["row"] = n
        facilities.setdefault(mtep_id(f["MTEP Project ID (Project) (Project)"]), []).append(f)
    for sheet in ("MTEP25 Data Pull - App. A", "MTEP25 Data Pull - App. B"):
        for n, c in sheet_rows(sheets[sheet], 1):
            out.append({"source": MISO_A, "artifact": art, "sheet": sheet, "row": n, "c": c, "as_of": as_of,
                        "native": mtep_id(c["MTEP Project ID"]), "facilities": facilities.get(mtep_id(c["MTEP Project ID"]), []),
                        "title": f"MISO MTEP25 Appendix A workbook ({sheet.split(' - ')[-1]}, as of {as_of})"})
    art = artifact_of(manifest, f"{MISO_EVAL}.xlsx")
    ws = workbook(cache / art["file"]).worksheets[0]
    stamp = next(ws.iter_rows(values_only=True, max_row=1))[5]
    as_of = stamp.date().isoformat() if isinstance(stamp, datetime) else None
    for n, c in sheet_rows(ws, 1):
        out.append({"source": MISO_EVAL, "artifact": art, "sheet": ws.title, "row": n, "c": c, "as_of": as_of,
                    "native": mtep_id(c["MTEP Project ID"]), "facilities": [],
                    "title": f"MISO MTEP Projects Under Evaluation (as of {as_of})"})
    return out


def build_miso(cache: Path, manifest: dict, osm: dict) -> tuple[list, list, dict]:
    linked = {p["native_id"]: p["_id"] for p in load_json(REPO_ROOT / "data" / "greatlakes" / "miso" / "projects.json")}
    groups: dict[str, list[dict]] = {}
    dispositions, as_of = [], {}
    for obs in miso_observations(cache, manifest):
        c = obs["c"]
        as_of[obs["source"]] = obs["as_of"]
        states = [s.strip() for s in str(c["State(s)"] or "").split(";") if s.strip()]
        locator = f"{obs['artifact']['file']}#{obs['sheet']}!row-{obs['row']}"
        where = {"source_id": f"southeast:{obs['source']}", "locator": locator,
                 "native_id": obs["native"], "name": c["Project Name"]}
        if not set(states) & set(STATES):
            # Out-of-scope rows keep only their locator and ID, which keeps the ledger small.
            dispositions.append({k: v for k, v in where.items() if k != "name"} | {
                "disposition": "excluded", "reason": f"outside F39 MISO states ({'; '.join(states) or 'none'})"})
        elif obs["native"] in linked:
            dispositions.append(where | {"disposition": "duplicate", "project_id": linked[obs["native"]],
                                         "reason": "same MTEP project ID already imported by F40 (Great Lakes)"})
        else:
            obs["states"], obs["where"] = states, where
            groups.setdefault(obs["native"], []).append(obs)
    projects = []
    for native, observations in groups.items():
        pid = f"southeast:miso-mtep:{native}"
        newest = observations[-1]
        c = newest["c"]
        for obs in observations:
            status = obs["c"]["Planning Status"] or ""
            obs |= {"group": STATUS.get(status[:2], "unknown"), "status": status, "date": day(obs["c"]["Expected ISD"]),
                    "date_field": "Expected ISD", "edition": slug(obs["source"]),
                    "evidence": evidence_of("miso", obs["artifact"], obs["where"]["locator"], obs["as_of"])}
            dispositions.append(obs["where"] | {"disposition": "accepted" if obs is newest else "duplicate",
                                                "project_id": pid,
                                                "reason": "F39 row of a MISO planning workbook" if obs is newest
                                                else "same MTEP project ID in a newer MISO workbook; kept as history"})
        name, description = c["Project Name"], c["Project Description"]
        keys = miso_keys(c["Submitting TO"] or "")
        facilities = [f for s in newest["states"] if s in osm for f in osm[s]]
        facility_rows = next((o["facilities"] for o in reversed(observations) if o["facilities"]), [])
        kv = {int(v) for v in (c.get("Max kV"), c.get("Min kV")) if isinstance(v, int | float) and v}
        kv |= {int(f["Max kV"]) for f in facility_rows if isinstance(f["Max kV"], int | float) and f["Max kV"]}
        # A program's facility row names one example site, not the program's place.
        named_subs = None if PROGRAM.search(name) else facility_names(facility_rows)
        if named_subs:
            center, candidate = locate_names(*named_subs, kv, facilities, keys, "facility_from_to")
        else:
            center, candidate = locate_text(name, description, facilities, keys)
        states = newest["states"]
        in_service = ({"raw": clean(c["Expected ISD"]), "value": newest["date"], "precision": "day"} if newest["date"]
                      else {"raw": None, "value": None, "precision": "unknown"})
        raw = {k: clean(v) for k, v in c.items() if k}
        raw["facilities"] = [{k: clean(v) for k, v in f.items() if k and v is not None} for f in facility_rows]
        projects.append({
            "_id": pid, "source_id": newest["where"]["source_id"], "native_id": native, "name": name,
            "description": description, "owner": c["Submitting TO"],
            "other_owners": sorted({o.strip() for f in facility_rows for o in str(f["Facility Owner(s)"] or "").split(";")
                                    if o.strip() and o.strip() != c["Submitting TO"]}),
            "planning_region": "miso", "states": [FIPS[s] for s in states if s in FIPS], "counties": [],
            "geography_basis": "source_state", "status": newest["status"], "status_group": newest["group"],
            "in_service": in_service, "center": center,
            "location_review": "unreviewed" if center else "unlocated", "location_candidate": candidate,
            "project_events": history(pid, native, observations),
            "estimated_cost": {"raw": c["Current Cost"],
                               "usd": c["Current Cost"] if isinstance(c["Current Cost"], int | float) else None},
            "evidence": {"page": None, "sheet": newest["sheet"], "row": newest["row"],
                         "source_sha256": newest["artifact"]["sha256"], "raw": raw},
        })
    return projects, dispositions, as_of


def spp_rows(path: Path) -> tuple[str, str | None, list[tuple[int, dict, dict]]]:
    """(sheet, zip member, [(row, normalized cells, original cells)]) of one Appendix 1; header drift fails closed."""
    member = None
    data: object = path
    if path.suffix == ".zip":
        archive = zipfile.ZipFile(path)
        members = [n for n in archive.namelist() if re.search(r"Appendix 1\.xls[xm]$", n)]
        if len(members) != 1:
            raise SystemExit(f"{path.name}: expected one Appendix 1 workbook, found {members}")
        member = members[0]
        data = io.BytesIO(archive.read(member))
    ws = workbook(data).worksheets[0]
    table = list(ws.iter_rows(values_only=True))
    def hkey(h) -> str:
        return re.sub(r"[\s_/]+", "", str(h or "")).lower()
    at = next(i for i, r in enumerate(table) if "uid" in [hkey(h) for h in r])
    header = [hkey(h) for h in table[at]]
    if missing := SPP_REQUIRED - set(header):
        raise SystemExit(f"{path.name}: Appendix 1 header changed, missing {sorted(missing)}")
    original = [" ".join(str(h).split()) if h is not None else "" for h in table[at]]
    rows = [(n, dict(zip(header, r, strict=False)), {k: clean(v) for k, v in zip(original, r, strict=False) if k})
            for n, r in enumerate(table[at + 1:], start=at + 2) if r[header.index("uid")] is not None]
    return ws.title, member, rows


def spp_title(edition: str) -> str:
    return f"SPP Quarterly Project Tracking Appendix 1, {edition[4:].upper()} {edition[:4]}"


def build_spp(cache: Path, manifest: dict, osm: dict) -> tuple[list, list]:
    listed: dict[str, list[dict]] = {}  # uid -> observations, newest edition first
    for edition in SPP_EDITIONS:
        art = artifact_of(manifest, spp_file(edition))
        sheet, member, rows = spp_rows(cache / art["file"])
        for n, c, raw in rows:
            uid = str(c["uid"]).strip()
            status = " ".join(str(c["projectstatus"] or "").split())
            locator = f"{art['file']}{'!' + member if member else ''}#{sheet}!row-{n}"
            listed.setdefault(uid, []).append({
                "edition": edition, "title": spp_title(edition), "row": n, "sheet": sheet, "c": c, "raw": raw,
                "artifact": art, "status": status, "group": spp_group(status),
                "date": day(c["projectownerindicatedin-servicedate"]),
                "date_field": "Project Owner Indicated In-Service Date",
                "where": {"source_id": f"southeast:spp-qpt-{edition}", "locator": locator, "native_id": uid,
                          "name": c["upgradename"]},
                "evidence": evidence_of("spp", art, locator, None)})
    projects, dispositions = [], []
    for uid, observations in listed.items():
        newest = observations[0]
        c = newest["c"]
        states = spp_states(c["state(s)"])
        f39 = [s for s in states if s in STATES]
        owner = str(c["projectowner"] or "").strip()
        pid = f"southeast:spp-qpt:{uid}"
        reason, project = None, None
        if not f39:
            reason = f"outside F39 states ({c['state(s)'] or 'none'})"
        else:
            rule, reason = spp_rule(owner, states)
            if rule != "excluded":
                project = spp_project(pid, uid, observations, states, f39, owner, osm)
                if rule == "border" and not project["center"]:
                    reason = (f"owner {owner or 'unassigned'} is not an established {'/'.join(f39)} transmission owner "
                              "and no facility the row names matches an OSM substation there")
                    project = None
        if project:
            projects.append(project)
        for obs in observations:
            newer = obs is not newest
            where = {k: v for k, v in obs["where"].items() if f39 or k != "name"}
            dispositions.append(where | (
                {"disposition": "duplicate" if newer else "accepted", "project_id": pid,
                 "reason": f"same UID in the newer {newest['title']}; kept as history" if newer
                 else "F39 upgrade in SPP's Quarterly Project Tracking"} if project else
                {"disposition": "excluded", "reason": (f"newest listing ({newest['title']}): " if newer else "") + reason}))
    return projects, dispositions


def spp_project(pid: str, uid: str, observations: list[dict], states: list[str], f39: list[str], owner: str,
                osm: dict) -> dict:
    newest = observations[0]
    c = newest["c"]
    name, description = " ".join(str(c["upgradename"]).split()), c.get("projectdescriptioncomments")
    keys = SPP_KEYS.get(owner, [])
    facilities = [f for s in f39 for f in osm[s]]
    center, candidate = locate_text(name, description, facilities, keys)
    if candidate["kind"] is None and (buses := bus_names(name, [c["frombusname"], c["tobusname"]])):
        kv = voltages_kv(name) | ({int(c["voltages(kv)"])} if isinstance(c.get("voltages(kv)"), int | float) else set())
        kind = "line" if len(buses) == 2 or re.search(r"\s[–-]\s", name) else "site"
        center, candidate = locate_names(kind, buses, kv, facilities, keys, "from_to_bus_names")
    in_service = ({"raw": str(clean(c["projectownerindicatedin-servicedate"])), "value": newest["date"],
                   "precision": "day"} if newest["date"] else {"raw": None, "value": None, "precision": "unknown"})
    return {
        "_id": pid, "source_id": newest["where"]["source_id"], "native_id": uid, "name": name,
        "description": description, "owner": None if owner.upper() in ("", "TBD") else owner, "other_owners": [],
        "planning_region": "spp", "states": [FIPS[s] for s in states], "counties": [],
        "geography_basis": "source_state", "status": newest["status"] or None, "status_group": newest["group"],
        "in_service": in_service, "center": center, "location_review": "unreviewed" if center else "unlocated",
        "location_candidate": candidate, "project_events": history(pid, uid, observations[::-1]),
        "evidence": {"page": None, "sheet": newest["sheet"], "row": newest["row"],
                     "source_sha256": newest["artifact"]["sha256"],
                     "raw": newest["raw"] | {"edition": newest["title"], "listed_in": [o["edition"] for o in observations]}},
    }


def build(cache: Path) -> dict:
    names = [f"{s}.xlsx" for s in MISO_FILES] + [spp_file(e) for e in SPP_EDITIONS] + \
        [f"osm-{s.lower()}.json" for s in STATES]
    manifest = verify_cache(cache, names)
    osm = load_osm(cache)
    miso, miso_disp, as_of = build_miso(cache, manifest, osm)
    spp, spp_disp = build_spp(cache, manifest, osm)
    projects = miso + spp
    count = Counter(p["source_id"] for p in projects)
    sources = []
    for source, (url, title) in MISO_FILES.items():
        art = manifest[f"{source}.xlsx"]
        sid = f"southeast:{source}"
        sources.append(source_record(sid, title, "miso", url, url, f"as of {as_of[source]}", art, projects,
                                     count[sid], [
            "F39 dense Southeast (C40): LA, AR, MS and KY rows. Locations are unreviewed OSM name candidates "
            "(C33 tiers, C38 operator guard). MISO's login-only project status report was not used."]))
    for edition in SPP_EDITIONS:
        art = manifest[spp_file(edition)]
        sid = f"southeast:spp-qpt-{edition}"
        sources.append(source_record(sid, spp_title(edition), "spp", SPP_INDEX, art["url"],
                                     f"{edition[4:].upper()} {edition[:4]}", art, projects, count[sid], [
            "F39 dense Southeast (C40): LA and AR upgrades. Older editions only add upgrades later editions "
            "dropped (closed-out upgrades leave the report) and each edition's changed in-service date as history.",
            "SPP state tags are filtered by an owner rule (see pipeline/southeast/misospp.py SPP_CORE/SPP_OUTSIDE)."]))
    return {"projects": projects, "sources": sources, "dispositions": miso_disp + spp_disp}


def source_record(sid, title, kind, landing, download, vintage, art, projects, n, notes) -> dict:
    return {"_id": sid, "title": title, "publisher": PLANNING[kind], "authority": "regional_planning_organization",
            "role": "project_plan", "landing_url": landing, "download_url": download, "publication_date": None,
            "vintage": vintage, "retrieved_at": art["retrieved_at"], "sha256": art["sha256"],
            "public_status": "verified_public", "import_status": "imported", "access_policy": "public_document",
            "planning_region": kind, "states": sorted({s for p in projects if p["source_id"] == sid for s in p["states"]}),
            "project_count": n, "notes": notes}



def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["fetch", "build"])
    parser.add_argument("--cache", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    if args.command == "fetch":
        fetch(args.cache)
        return 0
    result = build(args.cache)
    return write_batch(BATCH, result["projects"], result["sources"], result["dispositions"], args.check)


if __name__ == "__main__":
    sys.exit(main())

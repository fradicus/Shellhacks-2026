"""F45 Southwest: WestConnect TPPL projects (AZ/NM/CO) with C33/C38 candidates, plus WestTEC NV/UT lines (C42).

The WestConnect Transmission Plan Project List workbook is each sponsor's own entry: name, origin and termination
facilities, in-service year, development status and state. Its Origin and Termination are the project's endpoints;
they are matched by exact name to OSM substations in the row's state, as California was (C38). Nevada and Utah
have no TPPL rows, so a few WestTEC planned lines F42 did not import are placed from their own geometry.

From pipeline/:
  uv run python -m southwest.build fetch --cache /tmp/sw-f45
  uv run python -m southwest.build build --cache /tmp/sw-f45 [--check]
OSM data (c) OpenStreetMap contributors, ODbL 1.0.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from california.caiso import match, sha, slug
from common import REPO_ROOT, load_json, write_json
from greatlakes.match import candidate_center, facility_key, voltages_kv
from greatlakes.shared import fetch_into, fetch_osm, osm_extract, verify_cache
from texas.statewide import inside_geometry

OUT = REPO_ROOT / "data" / "southwest"
TPPL = "westconnect-tppl-2026-02"
TPPL_URL = "https://doc.westconnect.com/Documents.aspx?NID=21174&dl=1"
TPPL_LANDING = "https://regplanning.westconnect.com/tppl.htm"
WESTTEC = "westtec-10yr-planned-sw"
WESTTEC_URL = ("https://services3.arcgis.com/Iz3chmSt4P7oOoZy/arcgis/rest/services/WestTEC_10yr_Planned_Projects/"
               "FeatureServer/0/query?where=1%3D1&outFields=*&outSR=4326&f=geojson")
STATES_URL = ("https://tigerweb.geo.census.gov/arcgis/rest/services/TIGERweb/State_County/MapServer/0/query?"
              "where=STATE+IN+%28%2704%27%2C%2708%27%2C%2732%27%2C%2735%27%2C%2749%27%29&outFields=STATE%2CSTUSAB"
              "&outSR=4326&maxAllowableOffset=0.001&returnGeometry=true&f=geojson")
TPPL_STATES = {"Arizona": ("AZ", "04"), "New Mexico": ("NM", "35"), "Colorado": ("CO", "08")}
WESTTEC_STATES = {"32": "NV", "49": "UT"}
OPERATOR_KEYS = {
    "Arizona Public Service": ["ARIZONA PUBLIC SERVICE", "APS"],
    "Tucson Electric Power": ["TUCSON ELECTRIC"],
    "Public Service Company of New Mexico": ["PUBLIC SERVICE COMPANY OF NEW MEXICO", "PNM"],
    "El Paso Electric Company": ["EL PASO ELECTRIC"],
    "Public Service Company of Colorado/ Xcel Energy": ["XCEL", "PUBLIC SERVICE COMPANY OF COLORADO"],
    "Tri-State Generation and Transmission Association": ["TRI-STATE", "TRI STATE"],
    "Black Hills Energy": ["BLACK HILLS"],
}
STATUS = {"planned": "planned", "conceptual": "proposed", "in-service": "in_service",
          "under construction": "under_construction", "withdrawn": "cancelled"}
# Endpoint cells that describe a place instead of naming a facility.
NOT_A_NAME = re.compile(r"\b(TBD|N/?A|near|adjacent|existing|or|between|line|lines|area|point|POI|POCO|Sec|"
                        r"tap|location|future|proposed|undetermined|various|multiple)\b", re.I)
KV_CUT = re.compile(r"\s*\d+(?:\.\d+)?(?:\s*/\s*\d+(?:\.\d+)?)*\s*-?\s*kV\b.*$", re.I)
EQUIPMENT = re.compile(r"\s+(Substation|Susbtation|Switchyard|Switching Station|Station|Bus|Yard|Sub|Switch)$", re.I)
NAME = re.compile(r"[A-Z0-9][\w.'’-]*(?: [A-Za-z0-9][\w.'’-]*){0,4}")


def endpoint(cell) -> str | None:
    """The facility name an Origin/Termination cell states, or None when it states a description or nothing."""
    text = " ".join(str(cell or "").split())
    text = re.sub(r"\s*\([^)]*\)", "", text)  # "(formerly Hartt)"
    if not text or NOT_A_NAME.search(text):
        return None
    text = KV_CUT.sub("", text)
    while EQUIPMENT.search(text):
        text = EQUIPMENT.sub("", text)
    text = re.sub(r"^New\s+(?=\S+\s)", "", text).strip(" -")
    return text if NAME.fullmatch(text) and not text.isdigit() else None


def in_service(cell) -> dict:
    """Year precision exactly as the sponsor entered it; text such as TBD stays raw and unknown."""
    raw = None if cell is None else " ".join(str(cell).split())
    if isinstance(cell, int) and 2000 <= cell <= 2100:
        return {"raw": raw, "value": str(cell), "precision": "year"}
    return {"raw": raw, "value": None, "precision": "unknown"}


def locate(row: dict, facilities: list[dict]) -> tuple[dict | None, dict]:
    names = [endpoint(row["Origin"]), endpoint(row["Termination"])]
    keys = OPERATOR_KEYS.get(row["Sponsor"], [])
    kv = voltages_kv(row["Voltage"], str(row["Origin"] or ""), str(row["Termination"] or ""))
    if not any(names):
        kind, reason, matches = None, "no_named_endpoint", []
    else:
        same = names[0] and names[1] and facility_key(names[0]) == facility_key(names[1])
        kind, reason = ("site", None) if same or None in names else ("line", None)
        wanted = [n for n in names if n][:1] if kind == "site" else names
        matches = [match(n, facilities, keys, kv) for n in wanted]
    center = candidate_center(kind, matches) if kind else None
    found = [m for m in matches if m["status"] == "matched"]
    tier = None
    if center:
        tier = "candidate" if all(m["corroboration"] != ["unique_in_state"] for m in found) else "candidate_unique_name"
        center["evidence"] = center["evidence"].replace("corroborated by unique_in_state",
                                                        "the only facility with that name in the state (C33 name-only)")
    fields = ("id", "name", "operator", "voltage", "lat", "lon")
    endpoints = [{k: v for k, v in m.items() if k not in ("facility", "facility_ids")}
                 | ({"facility": {f: m["facility"].get(f) for f in fields}} if m["status"] == "matched" else {})
                 for m in matches]
    return center, {"rule": "C33", "tier": tier, "independent_review": False, "kind": kind, "reason": reason,
                    "voltages_kv": sorted(kv), "operator_keys": keys, "endpoints": endpoints,
                    "dataset": f"OpenStreetMap power=substation, {row['_usps']} (ODbL)"}


def read_tppl(path: Path) -> tuple[list[dict], str]:
    book = load_workbook(path, read_only=True, data_only=True)
    control = [c for r in book["Control"].iter_rows(values_only=True) for c in r if isinstance(c, datetime)]
    table = list(book["All Projects"].iter_rows(values_only=True))
    header = table[0]
    # The first header wins: a sort helper at the right edge repeats "projectid" over empty cells.
    first = {h: i for i, h in reversed(list(enumerate(header))) if h}
    rows = [{h: r[i] for h, i in first.items()} | {"_row": n} for n, r in enumerate(table[1:], start=2) if r[0] is not None]
    return rows, control[0].isoformat(sep=" ", timespec="minutes")


def tppl_project(row: dict, artifact: dict, facilities: list[dict]) -> dict:
    native = str(row["projectid"])
    pid = f"{TPPL}:{native}"
    usps, fips = TPPL_STATES[row["StateTraversed"]]
    status = row["Development"] or "Not stated"
    group = STATUS.get(status.lower(), "unknown")
    isd = in_service(row["InService"])
    center, candidate = locate(row | {"_usps": usps}, facilities)
    locator = f"sheet All Projects, row {row['_row']} (projectid {native})"
    evidence = {"publisher": "WestConnect", "url": artifact["url"], "artifact_sha256": artifact["sha256"],
                "locator": locator, "source_date": None, "retrieved_at": artifact["retrieved_at"],
                "access_review": "Public WestConnect TPPL workbook linked from the TPPL page; no login.", "facts": ""}
    events = []
    if group == "in_service":
        events.append({"id": f"{pid}:in-service", "type": "in_service", "date": isd["value"],
                       "precision": isd["precision"], "native_project_link": native,
                       "description": f"Development “{status}”; InService "
                                      + (f"{isd['value']}." if isd["value"] else "gives no year."),
                       "evidence": [evidence | {"facts": f"Development = {status}; InService = {isd['raw']}"}]})
    elif isd["value"]:
        events.append({"id": f"{pid}:planned-in-service", "type": "planned_milestone", "date": isd["value"],
                       "precision": "year", "native_project_link": native,
                       "description": f"Expected in service {isd['value']} (development status “{status}”).",
                       "evidence": [evidence | {"facts": f"InService = {isd['value']}"}]})
    modified = row.get("modifieddate")
    raw = {k: row.get(k) for k in ("projectid", "Sponsor", "ProjectName", "FacilityType", "Voltage", "Development",
                                   "Origin", "Termination", "InService", "SubRegionalID_Name", "StateTraversed")}
    return {
        "_id": pid, "source_id": TPPL, "native_id": native, "name": " ".join(str(row["ProjectName"]).split()),
        "description": (" ".join(str(row["Description"] or "").split())[:500] or None),
        "owner": row["Sponsor"], "other_owners": [], "planning_region": "westconnect", "states": [fips],
        "counties": [], "geography_basis": "source_state", "status": status, "status_group": group,
        "in_service": isd, "center": center, "location_review": "unreviewed" if center else "unlocated",
        "location_candidate": candidate, "project_events": events,
        "evidence": {"page": None, "sheet": "All Projects", "row": row["_row"], "source_sha256": artifact["sha256"],
                     "raw": raw | {"modifieddate": modified.date().isoformat() if modified else None}},
    }


def terminals(geometry: dict) -> list[list[float]]:
    parts = [geometry["coordinates"]] if geometry["type"] == "LineString" else geometry["coordinates"]
    return [parts[0][0], parts[-1][-1]]


def westtec_projects(raw: dict, states: dict, artifact: dict, published: set[str]) -> tuple[list[dict], list[dict]]:
    """Planned lines whose terminal-vertex mean lies in NV or UT; titles F42 already published are skipped."""
    titled: dict[str, list] = {}
    for f in raw["features"]:
        titled.setdefault(f["properties"]["Line_Name"], []).append(f)
    projects, dispositions = [], []
    for title, feats in sorted(titled.items()):
        native = slug(title)
        a, b = terminals(feats[0]["geometry"])[0], terminals(feats[-1]["geometry"])[1]
        lon, lat = round((a[0] + b[0]) / 2, 6), round((a[1] + b[1]) / 2, 6)
        fips = next((s for s, g in states.items() if s in WESTTEC_STATES and inside_geometry(lon, lat, g)), None)
        where = {"source_id": WESTTEC, "locator": f"WestTEC_10yr_Planned_Projects#{title}", "name": title}
        if native in published:
            dispositions.append(where | {"disposition": "duplicate", "reason": "published by F42"})
            continue
        if fips is None:
            dispositions.append(where | {"disposition": "excluded", "reason": "terminal mean outside NV/UT"})
            continue
        props = {k: v for k, v in feats[0]["properties"].items() if k not in ("FID",) and not k.startswith("Shape")}
        center = {"lat": lat, "lon": lon, "basis": "two",
                  "evidence": f"Official source geometry: WestTEC feature '{title}', mean of the line's terminal "
                              "vertices; not independently reviewed."}
        projects.append({
            "_id": f"{WESTTEC}:{native}", "source_id": WESTTEC, "native_id": native, "name": title,
            "description": None, "owner": None, "other_owners": [], "planning_region": "WestTEC 10-year",
            "states": [fips], "counties": [], "geography_basis": "source_geometry",
            "status": "; ".join(str(props[k]) for k in ("Upgrade_Ty",) if props.get(k)) or None,
            "status_group": "planned", "in_service": in_service(None), "center": center,
            "location_review": "unreviewed",
            "location_candidate": {"rule": "C33", "tier": "official", "independent_review": False, "kind": "line",
                                   "reason": None, "dataset": "WestTEC 10-year planned projects layer",
                                   "features": len(feats)},
            "project_events": [],
            "evidence": {"page": None, "sheet": "WestTEC_10yr_Planned_Projects", "row": None,
                         "source_sha256": artifact["sha256"], "raw": props | {"features": len(feats)}},
        })
        dispositions.append(where | {"disposition": "accepted", "_id": projects[-1]["_id"], "state": fips})
    return projects, dispositions


def fetch(cache: Path) -> None:
    cache.mkdir(parents=True, exist_ok=True)
    manifest = load_json(cache / "manifest.json") if (cache / "manifest.json").exists() else {}
    for name, url in (("tppl.xlsx", TPPL_URL), ("westtec.geojson", WESTTEC_URL), ("states.geojson", STATES_URL)):
        if name not in manifest:
            fetch_into(cache, name, url, manifest)
            write_json(cache / "manifest.json", manifest)
    for usps, _ in TPPL_STATES.values():
        if f"osm-{usps.lower()}.json" not in manifest:
            fetch_osm(cache, usps, manifest)
            write_json(cache / "manifest.json", manifest)


def build(cache: Path) -> dict[Path, object]:
    osm_files = [f"osm-{usps.lower()}.json" for usps, _ in TPPL_STATES.values()]
    manifest = verify_cache(cache, ["tppl.xlsx", "westtec.geojson", "states.geojson", *osm_files])
    facilities = {usps: osm_extract(json.loads((cache / f"osm-{usps.lower()}.json").read_bytes()), usps)
                  for usps, _ in TPPL_STATES.values()}
    rows, stamp = read_tppl(cache / "tppl.xlsx")
    projects, dispositions = [], []
    for row in rows:
        where = {"source_id": TPPL, "locator": f"All Projects row {row['_row']}", "name": row["ProjectName"]}
        state = row["StateTraversed"]
        if state not in TPPL_STATES:
            dispositions.append(where | {"disposition": "excluded", "reason": f"state {state}; C42 scope is AZ/NM/CO"})
            continue
        project = tppl_project(row, manifest["tppl.xlsx"], facilities[TPPL_STATES[state][0]])
        projects.append(project)
        dispositions.append(where | {"disposition": "accepted", "_id": project["_id"],
                                     "location": project["location_candidate"]["tier"] or "unlocated"})
    states = {f["properties"]["STATE"]: f["geometry"] for f in load_json(cache / "states.geojson")["features"]}
    published = {p["native_id"] for p in load_json(REPO_ROOT / "data" / "pnw" / "projects.json")
                 if p["source_id"] == "westtec-10yr-planned"}
    westtec, more = westtec_projects(load_json(cache / "westtec.geojson"), states, manifest["westtec.geojson"],
                                     published)
    projects = sorted(projects + westtec, key=lambda p: p["_id"])
    dispositions += more
    tppl, wt = manifest["tppl.xlsx"], manifest["westtec.geojson"]
    note = "F45 Southwest release (C42). Pins are unreviewed; none is independently confirmed. Source-bounded."
    sources = [
        {"_id": TPPL, "publisher": "WestConnect", "title": "WestConnect Transmission Plan Project List (TPPL) workbook",
         "authority": "regional_planning_organization", "role": "project_plan", "landing_url": TPPL_LANDING,
         "download_url": tppl["url"], "publication_date": None, "vintage": None,
         "retrieved_at": tppl["retrieved_at"], "sha256": tppl["sha256"], "public_status": "verified_public",
         "import_status": "imported", "access_policy": "public_document", "planning_region": "westconnect",
         "states": sorted({p["states"][0] for p in projects if p["source_id"] == TPPL}),
         "project_count": sum(p["source_id"] == TPPL for p in projects),
         "notes": [note, "Arizona, New Mexico and Colorado rows only. Candidate points are C33 exact-name OSM "
                   "matches with C38's operator guard. SRP, NV Energy and PacifiCorp do not file in the TPPL.",
                   f"The workbook's Control sheet TimeStamp reads {stamp}; its meaning (save or send) is not stated."]},
        {"_id": WESTTEC, "publisher": "WestTEC (layer hosted by BPA)",
         "title": "WestTEC 10-year planned projects layer (Nevada and Utah lines)",
         "authority": "regional_planning_organization", "role": "project_plan", "landing_url": WESTTEC_URL,
         "download_url": wt["url"], "publication_date": None, "vintage": None, "retrieved_at": wt["retrieved_at"],
         "sha256": wt["sha256"], "public_status": "verified_public", "import_status": "imported",
         "access_policy": "public_document", "planning_region": "WestTEC 10-year",
         "states": sorted({p["states"][0] for p in westtec}), "project_count": len(westtec),
         "notes": [note, "Only lines whose terminal-vertex mean is in Nevada or Utah and that F42 did not publish. "
                   "The layer gives no dates."]},
    ]
    return {OUT / "projects.json": projects, OUT / "sources.json": sources, OUT / "dispositions.json": dispositions,
            OUT / "osm-sources.json": {"publisher": "OpenStreetMap contributors", "rights": "ODbL 1.0; attribution "
                                       "required", "role": "candidate facility geometry only (C33)",
                                       "named_substations": {u: len(f) for u, f in facilities.items()},
                                       **{name: manifest[name] for name in osm_files}},
            OUT / "summary.json": summary(projects)}


def summary(projects: list[dict]) -> dict:
    located = [p for p in projects if p["center"]]
    return {"projects": len(projects), "located": len(located), "verified": 0,
            "distinct_points": len({(p["center"]["lat"], p["center"]["lon"]) for p in located}),
            "by_state": dict(sorted(Counter(p["states"][0] for p in projects).items())),
            "located_by_state": dict(sorted(Counter(p["states"][0] for p in located).items())),
            "by_tier": dict(sorted(Counter(p["location_candidate"]["tier"] for p in located).items())),
            "by_status": dict(sorted(Counter(p["status_group"] for p in projects).items())),
            "located_by_status": dict(sorted(Counter(p["status_group"] for p in located).items())),
            "located_with_dated_event": sum(any(e["date"] for e in p["project_events"]) for p in located),
            "located_not_in_service": sum(p["status_group"] not in {"in_service", "cancelled"} for p in located),
            "unlocated_reasons": dict(sorted(Counter(
                p["location_candidate"]["reason"] or ",".join(sorted({e["status"] for e in
                                                                      p["location_candidate"]["endpoints"]}))
                for p in projects if not p["center"]).items()))}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["fetch", "build"])
    parser.add_argument("--cache", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    if args.command == "fetch":
        fetch(args.cache)
        return 0
    from .publish import release

    outputs = build(args.cache)
    outputs |= release(outputs)
    if args.check:
        stale = [str(p) for p, v in outputs.items() if not p.exists() or load_json(p) != v]
        print("stale: " + ", ".join(stale) if stale else "ok")
        return 1 if stale else 0
    for path, value in outputs.items():
        write_json(path, value)
    print(json.dumps(outputs[OUT / "summary.json"], indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())

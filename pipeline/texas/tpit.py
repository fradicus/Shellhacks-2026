"""Extract a bounded public ERCOT TPIT sheet without publishing its raw workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SHEET = "PlannedTPIT071326NoCost"
SOURCE_ID = "ercot-tpit-2026-07"
GIS_URL = "https://gis.georgetowntexas.gov/arcgis/rest/services/GUS/GUSOPERATIONS_Webmap/MapServer/31"
CANDIDATE_ROWS = {67: "to", 105: "from", 285: "from", 333: "from"}


def _text(value: object) -> str | None:
    return str(value).strip() or None if value is not None else None


def _month(value: object) -> dict[str, str | None]:
    raw = value.isoformat() if isinstance(value, (date, datetime)) else _text(value)
    match = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])(?:-\d\d(?:T.*)?)?", raw or "")
    return {"raw": raw, "value": "-".join(match.groups()) if match else None,
            "precision": "month" if match else "unknown"}


def _facility(value: str | None) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\b(substation|sub|station)\b", "", value or "", flags=re.I)).strip().casefold()


def extract(workbook: Path, gis: Path) -> tuple[list[dict], dict]:
    source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    book = load_workbook(workbook, read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise ValueError(f"missing sheet {SHEET}")
    sheet = book[SHEET]
    headers = next(sheet.values)
    expected_title = (
        "JULY BASE FOR TRANSMISSION OWNER SUBMISSION FILE - "
        "TRANSMISSION PROJECT INFORMATION TRACKING (TPIT) FUTURE PROJECTS AS OF 7/13/2026"
    )
    if headers[0] != expected_title:
        raise ValueError("source vintage/header changed")
    headings = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    required_columns = [
        (0, "ERCOT Project Number"),
        (6, 'Transmission Status "under construction, planned or conceptual"'),
        (11, "Projected In-Service Date (Month/Yr)"),
        (18, "County Location for Substation or Starting Point for a Line"),
    ]
    for index, expected in required_columns:
        if headings[index] != expected:
            raise ValueError(f"column {index} changed")
    features = json.loads(gis.read_text())["features"]
    facilities: dict[str, list[dict]] = defaultdict(list)
    for feature in features:
        name = _facility(feature["properties"].get("Name"))
        facilities[name].append(feature)
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        if row[0] is None:
            continue
        native_id = _text(row[0])
        start = _text(row[4])
        end = _text(row[5])
        item = {
            "source_id": SOURCE_ID,
            "native_id": native_id,
            "sheet": SHEET,
            "row": row_number,
            "source_sha256": source_hash,
            "title": _text(row[1]),
            "terminal_from": start,
            "terminal_to": end,
            "status_raw": _text(row[6]),
            "sheet_cohort": "planned",
            "owner_raw": _text(row[8]),
            "projected_in_service": _month(row[11]),
            "actual_in_service": _month(row[12]),
            "voltage_kv_raw": _text(row[13]),
            "county_from_raw": _text(row[18]),
            "county_to_raw": _text(row[19]),
            "phase_raw": _text(row[29]),
            "disposition": "needs_scope_review",
            "candidate_location": None,
        }
        if row_number in CANDIDATE_ROWS:
            side = CANDIDATE_ROWS[row_number]
            terminal = start if side == "from" else end
            county = item["county_from_raw"] if side == "from" else item["county_to_raw"]
            possible = facilities[_facility(terminal)]
            if len(possible) != 1 or county != "Williamson":
                raise ValueError(f"candidate {row_number} lacks unique named facility/county")
            feature = possible[0]
            lon, lat = feature["geometry"]["coordinates"]
            if not (-106 <= lon <= -93 and 25 <= lat <= 37):
                raise ValueError(f"candidate {row_number} outside Texas bounds")
            item["candidate_location"] = {
                "tier": "candidate", "basis": "one", "lat": lat, "lon": lon,
                "facility_name": feature["properties"]["Name"],
                "facility_global_id": feature["properties"]["GlobalID"],
                "facility_owner_raw": feature["properties"].get("OWNER"),
                "gis_url": GIS_URL,
                "gis_crs": "EPSG:4326", "gis_sha256": hashlib.sha256(gis.read_bytes()).hexdigest(),
                "match": "unique normalized source terminal + named Williamson County",
                "independent_review": False,
            }
        records.append(item)
    ids = [item["native_id"] for item in records]
    if len(records) != 358 or len(set(ids)) != len(ids):
        raise ValueError("row count or ID uniqueness changed")
    if sum(item["candidate_location"] is not None for item in records) != len(CANDIDATE_ROWS):
        raise ValueError("candidate reconciliation failed")
    summary = {
        "source_id": SOURCE_ID, "source_sha256": source_hash, "sheet": SHEET,
        "observations": len(records), "unique_ids": len(set(ids)),
        "status_raw_counts": dict(sorted(Counter(item["status_raw"] or "unknown" for item in records).items())),
        "candidate_locations": len(CANDIDATE_ROWS), "official_locations": 0,
        "published_projects": 0, "scope_review_pending": len(records),
        "missing_titles": sum(item["title"] is None for item in records),
    }
    return records, summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("gis", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    records, summary = extract(args.workbook, args.gis)
    audit = json.loads((args.output / "source-audit.json").read_text())
    geo_audit = json.loads((args.output / "geo-source.json").read_text())
    if summary["source_sha256"] != audit["sha256"]:
        raise ValueError("ERCOT workbook differs from reviewed source")
    if hashlib.sha256(args.gis.read_bytes()).hexdigest() != geo_audit["sha256"]:
        raise ValueError("facility geometry differs from county-checked source")
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "planned-observations.json").write_text(json.dumps(records, indent=2) + "\n")
    (args.output / "planned-summary.json").write_text(json.dumps(summary, indent=2) + "\n")


if __name__ == "__main__":
    main()

"""F46 Midwest (IA, MO, KS, NE, ND, SD): SPP-approved upgrades with C33/C38 loose, labeled candidate locations.

SPP's Quarterly Project Tracking workbook lists every upgrade SPP has approved, its owner, status and the owner's
expected in-service date. Each upgrade is one record. Locations are candidates from named OSM substations in the
row's own states; none is independently reviewed (C43).

From pipeline/:
  uv run python -m midwest.build fetch --cache /tmp/midwest-f46    # SPP appendix zip + OSM substations, six states
  uv run python -m midwest.build build --cache /tmp/midwest-f46 [--check]
OSM data (c) OpenStreetMap contributors, ODbL 1.0.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

from california.caiso import match
from common import REPO_ROOT, load_json, write_json
from greatlakes.match import candidate_center, facilities_named, voltages_kv
from greatlakes.shared import fetch_into, fetch_osm, osm_extract, verify_cache

OUT = REPO_ROOT / "data" / "midwest"
STATES = {"IA": "19", "MO": "29", "KS": "20", "NE": "31", "ND": "38", "SD": "46"}
SOURCE_ID = "spp-qpt-2026q3"
ZIP = "spp-qpt-2026q3.zip"
URL = "https://www.spp.org/Documents/77416/" + quote("3Q 2026 Quarterly Project Tracking Report Appendix 1&2.zip")
LANDING = "https://www.spp.org/engineering/project-tracking-ntcs/"
MEMBER = "3Q 2026 Quarterly Project Tracking Report Appendix 1&2/Q3 2026 Quarterly Project Tracking Appendix 1.xlsm"
HEADER_ROW = 2  # row 1 is the report title
STATUS = {
    "planned": {"on schedule < 4", "on schedule > 4", "delay - mitigation", "delay - mitigation window"},
    "proposed": {"ntc - commitment window", "ntc-c project estimate window", "re-evaluation", "identified"},
    "in_service": {"complete", "in service", "closed out"},
    "cancelled": {"withdrawn"},
}
# OSM operator-name fragments for the SPP owner codes in these states. Owners without an entry get no operator
# corroboration and no operator guard.
OPERATOR_KEYS = {
    "BEPC": ["BASIN ELECTRIC"], "EREC": ["EAST RIVER"], "NPPD": ["NEBRASKA PUBLIC POWER", "NPPD"],
    "EKC": ["EVERGY", "WESTAR"], "EM": ["EVERGY", "KANSAS CITY POWER", "KCP&L"],
    "EMW": ["EVERGY", "KCP&L", "GREATER MISSOURI"], "WAPA": ["WESTERN AREA", "WAPA"], "EDE": ["EMPIRE", "LIBERTY"],
    "OPPD": ["OMAHA PUBLIC POWER", "OPPD"], "MIDW": ["MIDWEST ENERGY"], "MRES": ["MISSOURI RIVER"],
    "SEPC": ["SUNFLOWER"], "LES": ["LINCOLN ELECTRIC"], "TSGT": ["TRI-STATE"], "OGE": ["OKLAHOMA GAS", "OG&E"],
    "SPS": ["SOUTHWESTERN PUBLIC", "XCEL"], "CBPC": ["CORN BELT"], "ITCGP": ["ITC"], "NIPCO": ["NORTHWEST IOWA POWER"],
    "NEET": ["NEXTERA"],
}
# SPP name forms: "Craig 161 kV Ckt 2 Terminal Upgrade toward Midway 161 kV" is work at Craig; "toward" names the
# far end. "Wolf Creek 345kV Terminal Equipment", "Sweetwater 345kV GEN-2016-074 Interconnection" are sites.
TOWARD = re.compile(r"\s+toward\s+.*$", re.I)
SITE = re.compile(r"^(?P<name>.+?)\s*\d+(?:/\d+)*\s?kV?\b(?P<tail>.*)$", re.I)
SITE_WORK = re.compile(r"\b(Terminal|Interconnection|Substation|Sub|Transformer|Relay|Bus Tie|Shunt|Equipment)\b", re.I)
# Bus numbers, generator queue positions and border points are not facilities.
NOT_A_NAME = re.compile(r"^(?:S\s?\d+|Sub \d+|GEN-\d{4}-\d+|DISIS-\d{4}-\d+)$|\bBorder\b", re.I)
QUEUE = re.compile(r"^(?:GEN|DISIS)-\d{4}-\d+$", re.I)


def clean(value) -> str:
    return value.date().isoformat() if isinstance(value, datetime) else " ".join(str(value).split())


def status_group(status: str) -> str:
    text = " ".join(status.lower().split())
    return next((group for group, words in STATUS.items() if text in words), "unknown")


def row_states(cell) -> list[str]:
    return sorted({s.strip().upper() for s in re.split(r"[,/]", str(cell or "")) if s.strip()})


def named(upgrade: str) -> dict:
    """Facilities an SPP upgrade name states: F40's parser after SPP's own forms, then non-facility names dropped."""
    text = re.sub(r"(?<=[A-Za-z])(?=\d+\s?kV)", " ", upgrade)  # "Spring Creek345 kV"
    text = re.sub(r"\s+-\s+(?=\d+\s?kV)", " ", TOWARD.sub("", text))  # "Leland Olds - Finstad - 345 kV New Line"
    got = facilities_named(text, None)
    if got["reason"] == "no_named_facility" and (m := SITE.match(text)) and SITE_WORK.search(m["tail"]):
        got = {"kind": "site", "names": [m["name"].strip()], "from": "name", "reason": None}
    names = [None if n is None or NOT_A_NAME.search(n.strip()) else n for n in got["names"]]
    if got["kind"] and not any(names):
        return {"kind": None, "names": [], "from": "name", "reason": "not_a_facility"}
    if got["kind"] == "line" and any(n and QUEUE.match(n.strip()) for n in got["names"]):
        # "Holt County 345kV - GEN-2015-023 Addition": a generator's work at one site, not a line to a queue number.
        return got | {"kind": "site", "names": [n for n in names if n]}
    return got | {"names": names}


def locate(upgrade: str, kv_cell, states: list[str], facilities: dict[str, list[dict]], keys: list[str]
           ) -> tuple[dict | None, dict]:
    got = named(upgrade)
    pool = [f for s in states for f in facilities.get(s, [])]
    kv = voltages_kv(upgrade) | {int(v) for v in re.findall(r"\d+", str(kv_cell or "")) if int(v) > 0}
    matches = [match(n, pool, keys, kv) if n else {"status": "not_a_facility", "name": None} for n in got["names"]]
    center = candidate_center(got["kind"], matches) if got["kind"] else None
    fields = ("id", "name", "operator", "voltage", "state", "lat", "lon")
    endpoints = [{k: v for k, v in m.items() if k not in ("facility", "facility_ids")}
                 | ({"facility": {f: m["facility"].get(f) for f in fields}} if m["status"] == "matched" else {})
                 for m in matches]
    found = [m for m in matches if m["status"] == "matched"]
    tier = None
    if center:
        tier = "candidate" if all(m["corroboration"] != ["unique_in_state"] for m in found) else "candidate_unique_name"
    return center, {"rule": "C33", "tier": tier, "independent_review": False, "kind": got["kind"],
                    "reason": got["reason"], "voltages_kv": sorted(kv), "operator_keys": keys,
                    "endpoints": endpoints,
                    "dataset": f"OpenStreetMap power=substation, {'/'.join(states)} (ODbL)"}


def read_rows(zip_bytes: bytes) -> list[dict]:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archive:
        book = load_workbook(io.BytesIO(archive.read(MEMBER)), read_only=True, data_only=True)
    sheet = book.worksheets[0]
    table = list(sheet.iter_rows(values_only=True))
    header = [clean(h) if h is not None else "" for h in table[HEADER_ROW - 1]]
    rows = []
    for number, row in enumerate(table[HEADER_ROW:], start=HEADER_ROW + 1):
        cells = dict(zip(header, list(row) + [None] * (len(header) - len(row)), strict=True))
        if cells["UID"] is not None:
            rows.append(cells | {"_sheet": sheet.title, "_row": number})
    return rows


def project(row: dict, artifact: dict, facilities: dict[str, list[dict]]) -> dict:
    native = str(row["UID"])
    pid = f"{SOURCE_ID}:{native}"
    status = clean(row["Project Status"] or "Not stated")
    group = status_group(status)
    states = row_states(row["State(s)"])
    locator = f"{MEMBER.split('/')[-1]}, sheet {row['_sheet']}, row {row['_row']}"
    evidence = {"publisher": "Southwest Power Pool", "url": artifact["url"], "artifact_sha256": artifact["sha256"],
                "locator": locator, "source_date": None, "retrieved_at": artifact["retrieved_at"],
                "access_review": "Public SPP workbook; no login.", "facts": ""}
    cell = row["Project Owner Indicated In-Service Date"]
    value = cell.date().isoformat() if isinstance(cell, datetime) and cell.year >= 2000 else None
    precision = "day" if value else "unknown"
    events = []
    if group == "in_service":
        reported = value if value and value <= artifact["retrieved_at"][:10] else None
        events.append({"id": f"{pid}:status-in-service", "type": "in_service", "date": reported,
                       "precision": "day" if reported else "unknown", "native_project_link": native,
                       "description": (f"Project status “{status}”; owner-indicated in-service date {reported}."
                                       if reported else f"Project status “{status}”; the workbook gives no "
                                       "in-service date on or before retrieval."),
                       "evidence": [evidence | {"facts": f"Project Status = {status}"}]})
    elif value:
        events.append({"id": f"{pid}:owner-in-service", "type": "planned_milestone", "date": value,
                       "precision": "day", "native_project_link": native,
                       "description": f"Owner-indicated in-service date {value} (SPP Q3 2026 project tracking).",
                       "evidence": [evidence | {"facts": f"Project Owner Indicated In-Service Date = {value}"}]})
    owner = clean(row["ProjectOwner"])
    center, candidate = locate(clean(row["Upgrade Name"]), row["Voltages (kV)"], states, facilities,
                               OPERATOR_KEYS.get(owner, []))
    text = row["Project Description/ Comments"]
    raw_keys = ("NTC ID", "PID", "UID", "ProjectOwner", "State(s)", "Project Name", "Upgrade Name", "Project Type",
                "Project Owner Indicated In-Service Date", "Project Status", "Current Cost Estimate", "Voltages (kV)",
                "From Bus Name", "To Bus Name")
    return {
        "_id": pid, "source_id": SOURCE_ID, "native_id": native, "name": clean(row["Upgrade Name"]),
        "major_project": clean(row["Project Name"]) if row["Project Name"] else None, "part": clean(row["Upgrade Name"]),
        "description": clean(text)[:500] if text else None, "owner": owner, "other_owners": [], "planning_region": "spp",
        "states": [STATES[s] for s in states], "counties": [], "geography_basis": "source_state",
        "status": status, "status_group": group,
        "in_service": {"raw": clean(cell) if cell is not None else None, "value": value, "precision": precision},
        "center": center, "location_review": "unreviewed" if center else "unlocated",
        "location_candidate": candidate, "project_events": events,
        "evidence": {"page": None, "sheet": row["_sheet"], "row": row["_row"], "source_sha256": artifact["sha256"],
                     "raw": {k: clean(row[k]) if row[k] is not None else None for k in raw_keys}},
    }


def fetch(cache: Path) -> None:
    cache.mkdir(parents=True, exist_ok=True)
    manifest = load_json(cache / "manifest.json") if (cache / "manifest.json").exists() else {}
    if ZIP not in manifest:
        fetch_into(cache, ZIP, URL, manifest)
        write_json(cache / "manifest.json", manifest)
    for state in STATES:
        if f"osm-{state.lower()}.json" not in manifest:
            fetch_osm(cache, state, manifest)
            write_json(cache / "manifest.json", manifest)


def build(cache: Path) -> dict[Path, object]:
    osm_files = [f"osm-{s.lower()}.json" for s in STATES]
    manifest = verify_cache(cache, [ZIP, *osm_files])
    facilities = {s: osm_extract(json.loads((cache / f"osm-{s.lower()}.json").read_bytes()), s) for s in STATES}
    artifact = manifest[ZIP]
    projects, dispositions = [], []
    for row in read_rows((cache / ZIP).read_bytes()):
        states = row_states(row["State(s)"])
        where = {"source_id": SOURCE_ID, "sheet": row["_sheet"], "row": row["_row"], "uid": str(row["UID"]),
                 "name": clean(row["Upgrade Name"] or "")}
        if not states or not set(states) <= STATES.keys():
            reason = "no state listed" if not states else f"outside F46 states ({','.join(states)})"
            dispositions.append(where | {"disposition": "excluded", "reason": reason})
            continue
        record = project(row, artifact, facilities)
        projects.append(record)
        dispositions.append(where | {"disposition": "accepted", "_id": record["_id"],
                                     "location": record["location_candidate"]["tier"] or "unlocated"})
    if len({p["_id"] for p in projects}) != len(projects):
        raise SystemExit("SPP UID repeated within the workbook")
    projects.sort(key=lambda p: p["_id"])
    source = {
        "_id": SOURCE_ID, "publisher": "Southwest Power Pool",
        "title": "SPP Q3 2026 Quarterly Project Tracking Report, Appendix 1", "authority": "regional_planning_organization",
        "role": "project_plan", "landing_url": LANDING, "download_url": artifact["url"], "publication_date": None,
        "vintage": "2026 Q3", "retrieved_at": artifact["retrieved_at"], "sha256": artifact["sha256"],
        "public_status": "verified_public", "import_status": "imported", "access_policy": "public_document",
        "planning_region": "spp", "states": sorted({s for p in projects for s in p["states"]}),
        "project_count": len(projects),
        "notes": ["F46 Midwest release (C43): SPP upgrades whose listed states are all IA, MO, KS, NE, ND or SD. "
                  "Locations are unreviewed C33 candidates from named OSM substations; none is independently "
                  "confirmed. Source-bounded, not statewide coverage.",
                  f"sha256 is of the appendix zip; rows are read from its member “{MEMBER.split('/')[-1]}”."],
    }
    return {OUT / "projects.json": projects, OUT / "sources.json": [source], OUT / "dispositions.json": dispositions,
            OUT / "osm-sources.json": {"publisher": "OpenStreetMap contributors", "rights": "ODbL 1.0; attribution "
                                       "required", "role": "candidate facility geometry only (C33)",
                                       "named_substations": {s: len(f) for s, f in facilities.items()},
                                       "extracts": {name: manifest[name] for name in osm_files}},
            OUT / "summary.json": summary(projects)}


def summary(projects: list[dict]) -> dict:
    located = [p for p in projects if p["center"]]
    by_state = Counter(s for p in projects for s in p["states"])
    located_by_state = Counter(s for p in located for s in p["states"])
    return {"projects": len(projects), "located": len(located), "verified": 0,
            "distinct_points": len({(p["center"]["lat"], p["center"]["lon"]) for p in located}),
            "by_state": {s: {"projects": by_state[f], "located": located_by_state[f]} for s, f in STATES.items()},
            "by_tier": dict(sorted(Counter(p["location_candidate"]["tier"] for p in located).items())),
            "by_basis": dict(sorted(Counter(p["center"]["basis"] for p in located).items())),
            "by_status": dict(sorted(Counter(p["status_group"] for p in projects).items())),
            "located_by_status": dict(sorted(Counter(p["status_group"] for p in located).items())),
            "located_not_in_service": sum(p["status_group"] not in {"in_service", "cancelled"} for p in located),
            "located_with_dated_events": sum(any(e["date"] for e in p["project_events"]) for p in located),
            "unlocated_reasons": dict(sorted(Counter(
                p["location_candidate"]["reason"] or ",".join(sorted({e["status"] for e in
                                                                      p["location_candidate"]["endpoints"]}))
                for p in projects if not p["center"]).items()))}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["fetch", "build"])
    parser.add_argument("--cache", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    if args.command == "fetch":
        fetch(args.cache)
        return 0
    from .publish import release

    outputs = build(args.cache)
    outputs |= release(outputs)
    if args.check:
        stale = [str(p) for p, v in outputs.items() if not p.exists() or load_json(p) != v]
        print("stale: " + ", ".join(stale) if stale else "ok")
        return 1 if stale else 0
    for path, value in outputs.items():
        write_json(path, value)
    print(json.dumps(outputs[OUT / "summary.json"], indent=2))
    return 0


def sha(value: object) -> str:
    """Byte-identical to common.write_json, so the release can pin committed file hashes."""
    return hashlib.sha256((json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode()).hexdigest()


if __name__ == "__main__":
    sys.exit(main())
